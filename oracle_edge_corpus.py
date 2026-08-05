# -*- coding: utf-8 -*-
"""oracle_edge_corpus.py — ORACLE §3.5: สร้างคอร์ปัสขอบ แล้วตรวจว่าระบบ "ไม่เงียบ"

สร้างไฟล์ Excel ที่จำลองใบกำกับจริง (เลย์เอาต์เดียวกับ corpus) แต่จงใจมีข้อบกพร่อง
ที่ระบบต้องจับได้ แล้ววัดว่าแต่ละใบ "มี finding" หรือ "เงียบ".

เกณฑ์ (ปรัชญาระบบ): ใบที่ตรวจไม่ได้/ผิด ต้องไม่เงียบ — false-clean ร้ายแรงกว่า false-positive

รัน:  python3 oracle_edge_corpus.py [outdir]
"""
import io
import os
import sys
import glob
import shutil
import contextlib
import datetime as dt

from openpyxl import Workbook

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "_edge_corpus")

# คอลัมน์ตามเลย์เอาต์จริง (0-based): 1=seq 7=name/company 16=qty 17=unit 18=price 23=amount
C_SEQ, C_NAME, C_QTY, C_UNIT, C_PRICE, C_AMT = 1, 7, 16, 17, 18, 23
C_TAX, C_DATE, C_IV = 12, 18, 18


def put(ws, r, c, v):
    ws.cell(row=r + 1, column=c + 1, value=v)


def head(ws, iv, tax="0993000528883", company="กิจการร่วมค้า เอสเคที อินเทค-ซีทีซี (สำนักงานใหญ่)"):
    put(ws, 4, 8, "อ-043")
    put(ws, 4, C_TAX, float(tax))
    put(ws, 4, C_DATE, dt.datetime(2569, 5, 5))
    put(ws, 5, C_NAME, company)
    put(ws, 5, C_IV, iv)
    put(ws, 6, C_NAME, "33/4 ซอย อารีย์สัมพันธ์ 3 ถนนพหลโยธิน แขวงพญาไท ")
    put(ws, 7, C_NAME, "เขตพญาไท กรุงเทพมหานคร 10400")


def lines(ws, items, r0=12):
    for i, (name, qty, unit, price, amt) in enumerate(items):
        r = r0 + i
        put(ws, r, C_SEQ, float(i + 1))
        put(ws, r, C_NAME, name)
        put(ws, r, C_QTY, qty)
        put(ws, r, C_UNIT, unit)
        put(ws, r, C_PRICE, price)
        put(ws, r, C_AMT, amt)


def footer(ws, sub, vat, total, r0=27):
    put(ws, r0, C_AMT, sub)
    put(ws, r0 + 1, C_AMT, 0.0)
    put(ws, r0 + 3, 21, 0.0)
    put(ws, r0 + 3, C_AMT, sub)
    put(ws, r0 + 4, C_PRICE, 0.07)
    put(ws, r0 + 4, 21, vat)
    put(ws, r0 + 4, C_AMT, vat)
    put(ws, r0 + 5, 21, total)


GOOD = [("ท่อพีวีซี ฟ้า เอสซีจี ช้าง ชั้น 5 100 มม. 4 นิ้ว", 10.0, "เส้น", 300, 3000.0),
        ("ข้อต่อตรง พีวีซี 4 นิ้ว", 20.0, "ตัว", 50, 1000.0)]

CASES = {}


def case(name, build):
    CASES[name] = build


# 1) ใบปกติ (ควบคุม — ต้องเงียบ)
def _ok(ws):
    head(ws, "IV-69050901")
    lines(ws, GOOD)
    footer(ws, 4000.0, 280.0, 4280.0)


# 2) ใบที่ "ไม่มียอดท้ายบิลเลย" — เงินไม่เคยถูกอ่าน
def _nofooter(ws):
    head(ws, "IV-69050902")
    lines(ws, GOOD)


# 3) บรรทัดราคา 0 แต่มียอด (เคส ADR-193)
def _price0(ws):
    head(ws, "IV-69050903")
    lines(ws, [("ท่อพีวีซี ฟ้า เอสซีจี ช้าง ชั้น 5 100 มม. 4 นิ้ว", 10.0, "เส้น", 0, 3000.0),
               ("ข้อต่อตรง พีวีซี 4 นิ้ว", 20.0, "ตัว", 50, 1000.0)])
    footer(ws, 4000.0, 280.0, 4280.0)


# 4) VAT ไม่ใช่ 7%
def _badvat(ws):
    head(ws, "IV-69050904")
    lines(ws, GOOD)
    footer(ws, 4000.0, 123.45, 4123.45)


# 5) ผลรวมรายการ ≠ subtotal
def _badsum(ws):
    head(ws, "IV-69050905")
    lines(ws, GOOD)
    footer(ws, 9999.0, 699.93, 10698.93)


# 6) ใบซ้ำเลขเดิม ยอดเดิม (นับเงินซ้ำ)
def _dup(ws):
    head(ws, "IV-69050901")
    lines(ws, GOOD)
    footer(ws, 4000.0, 280.0, 4280.0)


# 7) ชีตว่าง
def _empty(ws):
    put(ws, 0, 0, None)


# 8) qty × price ≠ amount
def _qxp(ws):
    head(ws, "IV-69050908")
    lines(ws, [("ท่อพีวีซี ฟ้า เอสซีจี ช้าง ชั้น 5 100 มม. 4 นิ้ว", 10.0, "เส้น", 300, 2500.0),
               ("ข้อต่อตรง พีวีซี 4 นิ้ว", 20.0, "ตัว", 50, 1000.0)])
    footer(ws, 3500.0, 245.0, 3745.0)


case("ปกติ (ควบคุม)", _ok)
case("ไม่มียอดท้ายบิล", _nofooter)
case("ราคา=0 แต่มียอด", _price0)
case("VAT ไม่ใช่ 7%", _badvat)
case("ผลรวม≠subtotal", _badsum)
case("ใบซ้ำเลข+ยอดเดิม", _dup)
case("ชีตว่าง", _empty)
case("qty×price≠amount", _qxp)

if os.path.isdir(OUT):
    shutil.rmtree(OUT)
os.makedirs(OUT)

wb = Workbook()
wb.remove(wb.active)
sheet_of = {}
for i, (name, build) in enumerate(CASES.items(), start=1):
    ws = wb.create_sheet(title=str(i))
    build(ws)
    sheet_of[str(i)] = name
path = os.path.join(OUT, "EDGE_69_05.xlsx")
wb.save(path)
print(f"สร้างคอร์ปัสขอบ: {path}  ({len(CASES)} ชีต)")

# ── รันระบบจริงกับคอร์ปัสนี้ ────────────────────────────────────
import importlib
app = importlib.import_module("ปุ้มปุ้ย_ultimate_v9_modular")
with contextlib.redirect_stdout(io.StringIO()):
    app.reset_run_state()
    bills, fi = app.parse_all_files(sorted(glob.glob(os.path.join(OUT, "*.xls*"))))
    for b in bills:
        app.compute_bill_confidence(b)
    app.run_audit_core(bills, {}, isolate=True)

print(f"\nชีตทั้งหมด {len(CASES)} → ออกบิล {len(bills)}")
print("=" * 78)
print(f"{'ชีต':<22}{'เคส':<24}{'ยอด(sub/vat/total)':<26}{'รหัสที่ฟ้อง'}")
print("=" * 78)
seen = set()
silent = []
for b in sorted(bills, key=lambda x: str(x['sheet'])):
    sh = str(b['sheet'])
    seen.add(sh)
    codes = sorted({i['code'] for i in (b.get('issues') or []) if isinstance(i, dict)})
    amt = f"{b.get('subtotal')}/{b.get('vat')}/{b.get('total')}"
    print(f"{sh:<22}{sheet_of.get(sh,'?'):<24}{amt:<26}{','.join(codes) or '(เงียบ)'}")
    if not codes and sheet_of.get(sh) != "ปกติ (ควบคุม)":
        silent.append((sh, sheet_of.get(sh)))

missing = [(s, n) for s, n in sheet_of.items() if s not in seen]
print("=" * 78)
if missing:
    print("\n⚠ ชีตที่ไม่ออกบิลเลย (ต้องพิสูจน์ว่าไม่ใช่ใบกำกับจริง):")
    for s, n in missing:
        print(f"   ชีต {s}: {n}")
if silent:
    print("\n❌ ใบที่ 'ผิดแต่เงียบ' (false-clean):")
    for s, n in silent:
        print(f"   ชีต {s}: {n}")
else:
    print("\n✅ ไม่มีใบผิดที่เงียบ")
