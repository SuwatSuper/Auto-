# -*- coding: utf-8 -*-
"""build_consolidated_report.py — สร้าง Excel "สรุปข้อผิดพลาด (consolidated)" ที่ Agent ควรทำ

ยุบหลายรหัสที่ชี้ปัญหาเดียวกัน/จุดเดียวกัน → 1 ข้อสรุป แล้วแยกเป็น 3 เลน:
  • สรุปต้องแก้      — ปัญหาจริงที่ควรแก้ (ตัด master-artifact + ข้อสังเกตออกแล้ว)
  • ขึ้นกับ master    — ฟ้องเพราะ master ไม่ครบ (เอา master จริงมาใส่แล้วหาย)
  • ข้อสังเกต(review) — INFO/แนะนำ ไม่ใช่ must-fix

ใช้:
    PYTHONHASHSEED=0 PUOPUY_AUDIT_DATE=2026-06-02 \
      python3 build_consolidated_report.py [DATA_DIR] [OUT.xlsx]
อ่านอย่างเดียว (advisory) — ไม่แตะ engine/ผลตรวจ/golden hash.
"""
import warnings; warnings.filterwarnings("ignore")
import os
import sys
import glob
import io
import contextlib

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.chdir(os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DATA = sys.argv[1] if len(sys.argv) > 1 else "/mnt/project"
OUT = sys.argv[2] if len(sys.argv) > 2 else "/mnt/user-data/outputs/consolidated_errors.xlsx"
if not os.path.isdir(DATA):
    DATA = os.path.join("tests", "fixtures")

from golden_snapshot import MASTER, write_master_file
from report_precision import _xls_safe   # [ADR-119/R2] sanitize control char ก่อนเขียนเซลล์ (กัน IllegalCharacterError)
write_master_file("master_companies.json")
import importlib
app = importlib.import_module("ปุ้มปุ้ย_ultimate_v9_modular")
import issue_consolidator as IC

fl = sorted(glob.glob(os.path.join(DATA, "*.xls")) + glob.glob(os.path.join(DATA, "*.xlsx")))
with contextlib.redirect_stdout(io.StringIO()):
    app.reset_run_state()
    bills, fi = app.parse_all_files(fl)
    for b in bills:
        app.compute_bill_confidence(b)
    app.run_audit_core(bills, MASTER, isolate=True)

findings = IC.consolidate_all(bills)
stats = IC.summary_stats(bills)

# ── style ────────────────────────────────────────────────────────────────────
HF = PatternFill("solid", fgColor="1E40AF")
HFONT = Font(name="Arial", color="FFFFFF", bold=True, size=11)
CELL = Font(name="Arial", size=10)
SEV_FILL = {"CRITICAL": PatternFill("solid", fgColor="FEE2E2"),
            "ERROR": PatternFill("solid", fgColor="FEF3C7"),
            "WARNING": PatternFill("solid", fgColor="FEF9C3"),
            "INFO": PatternFill("solid", fgColor="EFF6FF")}
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)

COLS = [("ไฟล์", 20), ("ชีต", 8), ("เลขที่", 16), ("จุด", 12), ("หมวด", 16),
        ("สรุปปัญหา", 34), ("รหัสที่เกี่ยว", 24), ("#รหัส", 6),
        ("ความรุนแรง", 11), ("ตัวอย่างรายละเอียด", 50)]


def _write_sheet(wb, title, rows):
    ws = wb.create_sheet(title[:31])
    for ci, (h, w) in enumerate(COLS, 1):
        c = ws.cell(1, ci, h); c.fill = HF; c.font = HFONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[c.column_letter].width = w
    for ri, f in enumerate(rows, 2):
        vals = [f["file"], f["sheet"], f["iv"], f["spot"], f["category"],
                f["summary"], f["codes"], f["n_codes"], f["max_severity"], f["example"]]
        for ci, v in enumerate(vals, 1):
            # [ADR-119/R2] ค่ามาจาก findings ← parser ← Excel (อาจมี control char) → sanitize ก่อนเขียน
            #   มิฉะนั้น openpyxl โยน IllegalCharacterError ทำ consolidated_errors.xlsx เซฟไม่ออกทั้งไฟล์
            c = ws.cell(ri, ci, _xls_safe(v)); c.font = CELL; c.alignment = WRAP; c.border = BORDER
        ws.cell(ri, 9).fill = SEV_FILL.get(f["max_severity"], SEV_FILL["INFO"])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLS)).column_letter}{len(rows)+1}"
    return ws


wb = Workbook()
wb.remove(wb.active)

# ── ชีตสรุปภาพรวม (อ่านก่อน) ──────────────────────────────────────────────────
ov = wb.create_sheet("ภาพรวม")
ov.column_dimensions["A"].width = 38
ov.column_dimensions["B"].width = 14
ov_rows = [
    ("สรุปข้อผิดพลาด (Consolidated) — ยุบรหัสที่ชี้ปัญหาเดียวกัน", ""),
    (f"ข้อมูล: {os.path.basename(DATA)} — {len(fl)} ไฟล์ / {len(bills)} บิล", ""),
    ("", ""),
    ("issue ดิบ (กี่รหัสที่เด้ง)", stats["raw_issues"]),
    ("ยุบเป็น 'ปัญหาจริง' (consolidated)", stats["consolidated_findings"]),
    ("→ เฉพาะ 'ต้องแก้'", stats["fix_findings"]),
    ("   ใน (บิล)", stats["bills_with_fix"]),
    ("→ ขึ้นกับ master (เอา master จริงมาใส่แล้วหาย)", stats["by_bucket"].get("ขึ้นกับ master", 0)),
    ("→ ข้อสังเกต (review, ไม่ใช่ must-fix)", stats["by_bucket"].get("ข้อสังเกต (review)", 0)),
    ("", ""),
    ("— ปัญหาจริงต่อหมวด —", ""),
]
for cat, n in stats["by_category"].items():
    ov_rows.append((f"   {cat}", n))
for ri, (a, b) in enumerate(ov_rows, 1):
    ca = ov.cell(ri, 1, a); cb = ov.cell(ri, 2, b)
    ca.font = Font(name="Arial", size=11, bold=(ri <= 2 or a.startswith("—")))
    cb.font = Font(name="Arial", size=11, bold=True)
ov["A1"].fill = HF; ov["A1"].font = Font(name="Arial", color="FFFFFF", bold=True, size=12)

# ── 3 เลน ─────────────────────────────────────────────────────────────────────
fix = [f for f in findings if f["bucket"] == "ต้องแก้"]
mas = [f for f in findings if f["bucket"] == "ขึ้นกับ master"]
rev = [f for f in findings if f["bucket"] == "ข้อสังเกต (review)"]
_write_sheet(wb, f"สรุปต้องแก้ ({len(fix)})", fix)
_write_sheet(wb, f"ขึ้นกับ master ({len(mas)})", mas)
_write_sheet(wb, f"ข้อสังเกต ({len(rev)})", rev)

# [A1-FIX] OUT อาจเป็นชื่อไฟล์เปล่า (ไม่มีโฟลเดอร์นำ) → os.path.dirname คืน '' → makedirs('') ครัช.
#   สร้างโฟลเดอร์เฉพาะเมื่อมี dirname จริง.
_out_dir = os.path.dirname(OUT)
if _out_dir:
    os.makedirs(_out_dir, exist_ok=True)
wb.save(OUT)
print(f"✅ saved: {OUT}")
print(f"   ดิบ {stats['raw_issues']} → consolidated {stats['consolidated_findings']} "
      f"(ต้องแก้ {len(fix)} / master {len(mas)} / review {len(rev)})")
