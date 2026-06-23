# -*- coding: utf-8 -*-
"""test_recheck_20260620.py — ตรึงทุกการแก้จากรอบ re-audit 2026-06-20.

รันแบบ standalone:  python3 test_recheck_20260620.py   (exit 0 = ผ่าน)
ครอบ: C1(golden ดูที่ regression), F1/F2, C3/A-C1, C4/REP-C1, REP-C2, REP-M1,
      L1/L2/L3, L4, A-M1, A-L3, P3, P4. ทุกเคส golden-neutral (ไม่แตะผลตรวจข้อมูลจริง).
"""
import os
import sys
import math
import tempfile
from decimal import Decimal

os.environ.setdefault("PYTHONHASHSEED", "0")
os.environ.setdefault("PUOPUY_AUDIT_DATE", "2026-06-02")
os.environ.setdefault("PUOPUY_ALLOW_VERSION_MISMATCH", "1")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

_passed = 0
_failed = 0


def check(cond, label):
    global _passed, _failed
    if cond:
        _passed += 1
        print(f"  ✅ {label}")
    else:
        _failed += 1
        print(f"  ❌ {label}")


print("\n[F1] run_rules: บิล item ขาดคีย์ → ไม่ข้ามกฎเงียบ (ITM/VAT001 ต้องรัน)")
import state
import rules_engine
if hasattr(state, "_SYSTEM_ISSUES"):
    state._SYSTEM_ISSUES.clear()
bill = {"company": "บริษัท เอ จำกัด", "items": [{"name": "ของ", "seq": 1}],
        "subtotal": 100.0, "vat": 7.0, "total": 107.0}
rules_engine.run_rules(bill, {}, {"month": None})
_skipped = {i.get("code") for i in getattr(state, "_SYSTEM_ISSUES", [])}
check("SYS-VAT001" not in _skipped and "SYS-ITM001" not in _skipped,
      "ITM001/VAT001 ไม่ถูกข้ามเป็น SYS-* (item ถูก setdefault ครบ 7 คีย์)")

print("\n[F2] r_vat001: amount แปลงเป็นเลขไม่ได้ → ไม่ครัช (VAT001 ไม่ถูกข้ามเงียบ)")
from rules_engine_rules_b import r_vat001
def _bill_amt(a):
    return {"items": [{"name": "a", "seq": 1, "qty": 1.0, "price": 10.0, "amount": a, "unit": ""}],
            "subtotal": 10.0, "vat": None, "total": None}
ok_f2 = True
for bad in ("abc", True, None):
    try:
        r_vat001(_bill_amt(bad), None, {})
    except Exception:
        ok_f2 = False
check(ok_f2, "r_vat001 ทน amount = 'abc'/True/None (กรองหลัง _D เหมือน vat006/007)")

print("\n[L1] VAT rules: quantize ยอดมหึมา → [] ไม่ครัช")
from rules_engine_rules_b import r_vat002, r_vat003
from rules_engine_rules_c import r_vat006, r_vat007
big = {"subtotal": 1e300, "vat": 1e290, "total": 1e305,
       "items": [{"amount": 1e300, "seq": 1}]}
ok_l1 = True
for fn in (r_vat002, r_vat003, r_vat006, r_vat007):
    try:
        fn(big, None, {})
    except Exception:
        ok_l1 = False
check(ok_l1, "r_vat002/003/006/007 ห่อ quantize (InvalidOperation) แล้ว")

print("\n[L2] r_itm002: seq มหึมา → ไม่สร้าง set ยักษ์ (กัน DoS)")
from rules_engine_rules_a import r_itm002
import time as _t
_t0 = _t.time()
_out = r_itm002({"items": [{"seq": 1}, {"seq": 9_000_000}]}, None, {})
check((_t.time() - _t0) < 1.0 and any("ผิดช่วง" in x for x in _out),
      "seq 9,000,000 → ฟ้องผิดช่วงเร็ว (<1s) ไม่สร้าง range ยักษ์")

print("\n[L3] r_dt004: เดือน > 12 (date-like) → ฟ้อง")
from rules_engine_rules_c import r_dt004
class _D13:
    year, month, day = 2025, 13, 5
check(any("เดือน 13" in x for x in r_dt004({"iv_date": _D13()}, None, {})),
      "month=13 ถูกตรวจจับ (real datetime สร้างไม่ได้ → golden-neutral)")

print("\n[P3] _cell_to_num: เลขติดลบบัญชี (1,234.50) → -1234.50")
from parser_p0a import _cell_to_num
check(_cell_to_num("(1,234.50)") == -1234.50 and _cell_to_num("(500)") == -500.0
      and _cell_to_num("500") == 500.0 and _cell_to_num("(abc)") is None,
      "วงเล็บตัวเลขล้วน → ลบ; ปกติ/ข้อความ → เหมือนเดิม")

print("\n[P4] parse_date_any: วันที่ตัวเลขมี label นำหน้า → parse ได้")
from puopuy_dates import parse_date_any
_d1 = parse_date_any("วันที่ 11/05/69")
_d2 = parse_date_any("Date: 5/5/69")
_d3 = parse_date_any("5/5/69")
check(_d1 is not None and _d1.year == 2026 and _d2 is not None and _d3 is not None,
      "'วันที่ 11/05/69' / 'Date: 5/5/69' / '5/5/69' → ปี 2026 (ค.ศ.)")

print("\n[A-C1] _q2: ยอดมหึมา → ไม่ครัช (VerificationAgent ไม่ทิ้งผลโหวตทั้งชุด)")
from agents.verification_lenses_base import _q2
ok_ac1 = True
try:
    _q2(Decimal("1e30"))
except Exception:
    ok_ac1 = False
check(ok_ac1, "_q2(Decimal('1e30')) ไม่ raise InvalidOperation")

print("\n[A-M1] ConfidenceAgent: key low-conf ด้วย iv_number_raw (โบนัสไม่ตาย)")
import inspect
import agents.confidence_agent as _ca
_src = inspect.getsource(_ca.ConfidenceAgent._run)
check("iv_number_raw" in _src, "low_conf_bills key ใช้ iv_number_raw (ตรง bill_ref/f.iv)")

print("\n[A-L3] verify_severities: ส่งสตริงเดี่ยว → ไม่ถูกแตกเป็นตัวอักษร")
_src2 = inspect.getsource(__import__("agents.verification_agent", fromlist=["x"]).VerificationAgent._run)
check("isinstance(_sev, str)" in _src2, "สตริงเดี่ยว ('ERROR') ถูกห่อเป็น tuple 1 สมาชิก")

print("\n[REP-M1] analytics._num: ±inf/NaN → 0.0")
from analytics import _num
check(_num(float("inf")) == 0.0 and _num(float("-inf")) == 0.0
      and _num(float("nan")) == 0.0 and _num(5.5) == 5.5,
      "_num กัน NaN และ ±inf (เหลือค่าจำกัดเท่าเดิม)")

print("\n[L4] _df_safe: control-char + NaN ใน DataFrame → เขียน to_excel ได้")
import pandas as pd
import openpyxl
from reporting_p0 import _df_safe
_df = pd.DataFrame([{"a": "X\x07Y", "b": float("nan"), "c": 1000.0}])
_p = os.path.join(tempfile.mkdtemp(), "t.xlsx")
_ok_l4 = True
try:
    _df_safe(_df).to_excel(_p, index=False)
    _vals = [c.value for r in openpyxl.load_workbook(_p).active.iter_rows() for c in r]
    _ok_l4 = "\x07" not in str(_vals) and not any(isinstance(v, float) and v != v for v in _vals)
except Exception:
    _ok_l4 = False
check(_ok_l4, "full-mode export ทน control-char + NaN")

print("\n[C4/REP-C1 + REP-C2] build_clean_report: master_key มี \\x07 + บิล NaN → ออกไฟล์ ไม่มีเซลล์ NaN")
from reporting_p2 import build_clean_report
def _mk(company, sub):
    b = {"file": "F.xls", "sheet": "1", "company": company, "iv_number": "IV1",
         "iv_number_raw": "IV1", "iv_date": None, "subtotal": sub,
         "vat": (None if sub != sub else 7.0), "total": (sub if sub != sub else sub + 7.0),
         "items": [{"seq": 1, "name": "x", "name_raw": "x", "qty": 1.0, "unit": "",
                    "price": (None if sub != sub else sub), "amount": sub}],
         "issues": [{"code": "VAT001", "severity": "CRITICAL", "category": "VAT", "name": "t", "detail": "d"}]}
    rules_engine.run_rules(b, {}, {})
    b["master_key"] = company
    return b
_nan = float("nan")
_b1 = _mk("ACME\x07CORP", 1000.0)
_o1 = os.path.join(tempfile.mkdtemp(), "r1.xlsx")
_ok_c4 = build_clean_report([_b1], [{"key": "ACME\x07CORP", "bills": [_b1]}], [], [], [], _o1)
check(_ok_c4 and os.path.exists(_o1), "control-char master_key → รายงานออกไฟล์ (ไม่คืน False)")

_b2, _b3, _b4 = _mk("GOOD", 1000.0), _mk("BAD", _nan), _mk("GOOD2", 500.0)
_o2 = os.path.join(tempfile.mkdtemp(), "r2.xlsx")
_ok_c2 = build_clean_report([_b2, _b3, _b4],
                            [{"key": "GOOD", "bills": [_b2]}, {"key": "BAD", "bills": [_b3]},
                             {"key": "GOOD2", "bills": [_b4]}], [], [], [], _o2)
_nan_cells = 0
if _ok_c2:
    _wb = openpyxl.load_workbook(_o2)
    _nan_cells = sum(1 for ws in _wb for row in ws.iter_rows() for c in row
                     if isinstance(c.value, float) and c.value != c.value)
check(_ok_c2 and _nan_cells == 0, "บิล NaN → รายงานออก + ไม่มีเซลล์ NaN (ยอดไม่หาย)")

print("\n" + "=" * 64)
print(f"RECHECK 2026-06-20: ผ่าน {_passed} / ล้มเหลว {_failed}")
print("=" * 64)
if _failed:
    print("RESULT: ❌ พบ regression")
    sys.exit(1)
print("RESULT: ✅ ทุกการแก้ถูกตรึงครบ")
sys.exit(0)
