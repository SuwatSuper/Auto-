# -*- coding: utf-8 -*-
"""test_bughunt_hardening.py — ด่านกันถอยหลังของการล่าบั๊ก v9.3.1 (รอบ hardening)

ครอบคลุมบั๊ก "golden-safe" ที่ reproduce ได้ (ทาง error/edge/malformed เท่านั้น —
ข้อมูลจริงในชุด golden ไม่เดินทางนี้ จึง hash ไม่ขยับ):

  • Parse-S1 : _dic_int_run ระเบิด OverflowError เมื่อเซลล์เป็นข้อความ 'inf'/'1e400'
               → detect คอลัมน์ครัช → parse_file ดักที่ระดับชีต → บิล "ทั้งชีต" หายเงียบ
  (รอบถัดไปจะ append: Parse-M1/M3 (Decimal/inf), Master-S1, Master-M1/M2)

รันเดี่ยว:  PYTHONHASHSEED=0 PUOPUY_AUDIT_DATE=2026-06-02 python3 test_bughunt_hardening.py
exit 0 = ผ่านหมด, 1 = พบ regression
"""
import os
import sys
import io
import shutil
import tempfile
import contextlib
import importlib

import numpy as np
from openpyxl import load_workbook

_buf = io.StringIO()
with contextlib.redirect_stdout(_buf):
    importlib.import_module('ปุ้มปุ้ย_ultimate_v9_modular')
import parser as P
import parser_p0a as P0A

PASS, FAIL = 0, []


def check(cond, label):
    global PASS
    if cond:
        PASS += 1
        print(f"  ✅ {label}")
    else:
        FAIL.append(label)
        print(f"  ❌ {label}")


def _quiet(fn):
    """รัน fn() (ดูด stdout) — คืน (ok, result). ok=False ถ้า throw."""
    try:
        b = io.StringIO()
        with contextlib.redirect_stdout(b):
            return True, fn()
    except BaseException as e:        # parser ห้ามล้มทุกชนิด (รวม SystemExit)
        return False, f"{type(e).__name__}: {str(e)[:120]}"


FIXTURE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       'tests', 'fixtures', 'fixture_invoices.xlsx')
TMP = tempfile.mkdtemp(prefix="pukpui_harden_")


def _inject_cell(value, where=('D', 14)):
    """copy fixture แล้วยัดค่า pathological ลง 1 เซลล์ของชีตแรก → คืน path ไฟล์ใหม่."""
    out = os.path.join(TMP, f"inj_{abs(hash((str(value), where)))}.xlsx")
    shutil.copy2(FIXTURE, out)
    wb = load_workbook(out)
    ws = wb.worksheets[0]
    ws[f"{where[0]}{where[1]}"] = value
    wb.save(out)
    wb.close()
    return out


print("=" * 64)
print("BUGHUNT HARDENING v9.3.1 — golden-safe crash/data-integrity guards")
print("=" * 64)

# ── baseline: fixture ปกติต้องได้บิล ≥ 1 (กันเทสเองพัง) ──────────────────────────
ok0, base_bills = _quiet(lambda: P.parse_file(FIXTURE))
check(ok0 and isinstance(base_bills, list) and len(base_bills) >= 1,
      f"baseline: fixture ปกติ parse ได้ {len(base_bills) if ok0 else 'CRASH'} บิล")

# ════════════════════════════════════════════════════════════════════════════
# Parse-S1 — _dic_int_run OverflowError ('inf'/'1e400')
# ════════════════════════════════════════════════════════════════════════════
print("\n[Parse-S1] เซลล์ข้อความ 'inf'/'1e400' ต้องไม่ทำบิลทั้งชีตหาย")

# unit: _dic_int_run ไม่ครัช + ไม่นับ inf เป็นลำดับสินค้า
for bad in ('inf', '-inf', 'Infinity', '1e400'):
    M = np.array([[bad], [5], [3]], dtype=object)
    ok, res = _quiet(lambda: P0A._dic_int_run(M, 0))
    check(ok and res == [5, 3], f"_dic_int_run ข้าม {bad!r} อย่างปลอดภัย (ได้ {res!r})")

# end-to-end: ไฟล์ที่มีเซลล์ 'inf' ต้องยัง parse บิลได้ (เดิม = 0 บิล + ครัชระดับชีต)
for bad in ('inf', '1e400'):
    f = _inject_cell(bad)
    ok, bills = _quiet(lambda: P.parse_file(f))
    n = len(bills) if ok and isinstance(bills, list) else -1
    check(ok and n >= 1, f"parse_file(ไฟล์มีเซลล์ {bad!r}) คืน {n} บิล (ต้อง ≥1, ไม่ครัช)")

# ════════════════════════════════════════════════════════════════════════════
shutil.rmtree(TMP, ignore_errors=True)
print("\n" + "=" * 64)
print(f"BUGHUNT HARDENING: ผ่าน {PASS} / ล้มเหลว {len(FAIL)}")
print("=" * 64)
if FAIL:
    for f in FAIL:
        print(f"  • {f}")
    print("RESULT: ❌ พบ regression")
    sys.exit(1)
print("RESULT: ✅ การ์ด hardening ครบ ไม่มี regression")
sys.exit(0)
