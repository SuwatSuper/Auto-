# -*- coding: utf-8 -*-
"""test_bughunt_hardening.py — ด่านกันถอยหลังของการล่าบั๊ก v9.3.1 (รอบ hardening)

ครอบคลุมบั๊ก "golden-safe" ที่ reproduce ได้ (ทาง error/edge/malformed เท่านั้น —
ข้อมูลจริงในชุด golden ไม่เดินทางนี้ จึง hash ไม่ขยับ):

  • Parse-S1 : _dic_int_run ระเบิด OverflowError เมื่อเซลล์เป็นข้อความ 'inf'/'1e400'
               → detect คอลัมน์ครัช → parse_file ดักที่ระดับชีต → บิล "ทั้งชีต" หายเงียบ
  • Parse-M1 : ยอด subtotal มหึมา → Decimal.quantize ระเบิด InvalidOperation → บิลหายทั้งชีต/ไฟล์
  • Parse-M3 : _cell_to_num ปล่อย ±inf หลุด (f!=f จับแค่ NaN) → คูณ Decimal ระเบิดปลายน้ำ
  • Master-S1: write_master_file kill-before-atexit + ทับ .user.bak → master จริง+backup หายถาวร
  • Master-M1/M2: save_master เขียนไม่ atomic (kill = master ครึ่ง/ว่าง) + .bak ถูก save ที่หดทับ

รันเดี่ยว:  PYTHONHASHSEED=0 PUOPUY_AUDIT_DATE=2026-06-02 python3 test_bughunt_hardening.py
exit 0 = ผ่านหมด, 1 = พบ regression
"""
import os
import sys
import io
import shutil
import tempfile
import contextlib
import importlib

import numpy as np
from openpyxl import load_workbook

_buf = io.StringIO()
with contextlib.redirect_stdout(_buf):
    importlib.import_module('ปุ้มปุ้ย_ultimate_v9_modular')
import json
import parser as P
import parser_p0a as P0A
import parser_p2 as P2
import golden_snapshot as GS
import master as mstr

PASS, FAIL = 0, []


def check(cond, label):
    global PASS
    if cond:
        PASS += 1
        print(f"  ✅ {label}")
    else:
        FAIL.append(label)
        print(f"  ❌ {label}")


def _quiet(fn):
    """รัน fn() (ดูด stdout) — คืน (ok, result). ok=False ถ้า throw."""
    try:
        b = io.StringIO()
        with contextlib.redirect_stdout(b):
            return True, fn()
    except BaseException as e:        # parser ห้ามล้มทุกชนิด (รวม SystemExit)
        return False, f"{type(e).__name__}: {str(e)[:120]}"


FIXTURE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       'tests', 'fixtures', 'fixture_invoices.xlsx')
TMP = tempfile.mkdtemp(prefix="pukpui_harden_")


def _inject_cell(value, where=('D', 14)):
    """copy fixture แล้วยัดค่า pathological ลง 1 เซลล์ของชีตแรก → คืน path ไฟล์ใหม่."""
    out = os.path.join(TMP, f"inj_{abs(hash((str(value), where)))}.xlsx")
    shutil.copy2(FIXTURE, out)
    wb = load_workbook(out)
    ws = wb.worksheets[0]
    ws[f"{where[0]}{where[1]}"] = value
    wb.save(out)
    wb.close()
    return out


print("=" * 64)
print("BUGHUNT HARDENING v9.3.1 — golden-safe crash/data-integrity guards")
print("=" * 64)

# ── baseline: fixture ปกติต้องได้บิล ≥ 1 (กันเทสเองพัง) ──────────────────────────
ok0, base_bills = _quiet(lambda: P.parse_file(FIXTURE))
check(ok0 and isinstance(base_bills, list) and len(base_bills) >= 1,
      f"baseline: fixture ปกติ parse ได้ {len(base_bills) if ok0 else 'CRASH'} บิล")

# ════════════════════════════════════════════════════════════════════════════
# Parse-S1 — _dic_int_run OverflowError ('inf'/'1e400')
# ════════════════════════════════════════════════════════════════════════════
print("\n[Parse-S1] เซลล์ข้อความ 'inf'/'1e400' ต้องไม่ทำบิลทั้งชีตหาย")

# unit: _dic_int_run ไม่ครัช + ไม่นับ inf เป็นลำดับสินค้า
for bad in ('inf', '-inf', 'Infinity', '1e400'):
    M = np.array([[bad], [5], [3]], dtype=object)
    ok, res = _quiet(lambda: P0A._dic_int_run(M, 0))
    check(ok and res == [5, 3], f"_dic_int_run ข้าม {bad!r} อย่างปลอดภัย (ได้ {res!r})")

# end-to-end: ไฟล์ที่มีเซลล์ 'inf' ต้องยัง parse บิลได้ (เดิม = 0 บิล + ครัชระดับชีต)
for bad in ('inf', '1e400'):
    f = _inject_cell(bad)
    ok, bills = _quiet(lambda: P.parse_file(f))
    n = len(bills) if ok and isinstance(bills, list) else -1
    check(ok and n >= 1, f"parse_file(ไฟล์มีเซลล์ {bad!r}) คืน {n} บิล (ต้อง ≥1, ไม่ครัช)")

# ════════════════════════════════════════════════════════════════════════════
# Parse-M1 — Decimal.quantize InvalidOperation (ยอดมหึมา) ใน _pb_finalize_amounts
# ════════════════════════════════════════════════════════════════════════════
print("\n[Parse-M1] ยอด subtotal มหึมา (≥~1e28) ต้องไม่ทำ derive-VAT ครัช")
for sub in (1e30, 1e29, 9.9e27):
    r = {'subtotal': sub, 'vat': None, 'total': None, 'items': []}
    ok, _ = _quiet(lambda: P2._pb_finalize_amounts(r))
    check(ok, f"_pb_finalize_amounts(subtotal={sub:.0e}) ไม่ครัช (vat={r.get('vat')})")
# พฤติกรรมปกติต้องเหมือนเดิมเป๊ะ: 2500 → vat 175.0, total 2675.0
r_ok = {'subtotal': 2500.0, 'vat': None, 'total': None, 'items': []}
_quiet(lambda: P2._pb_finalize_amounts(r_ok))
check(r_ok.get('vat') == 175.0 and r_ok.get('total') == 2675.0,
      f"ยอดปกติ 2500 → derive vat={r_ok.get('vat')} total={r_ok.get('total')} (ต้อง 175.0/2675.0)")

# ════════════════════════════════════════════════════════════════════════════
# Parse-M3 — _cell_to_num ต้องกรอง ±inf/NaN (ไม่ปล่อยหลุดไปคูณ Decimal ปลายน้ำ)
# ════════════════════════════════════════════════════════════════════════════
print("\n[Parse-M3] _cell_to_num ต้องคืน None สำหรับ NaN/±inf และคงค่าจริง")
check(P0A._cell_to_num(float('inf')) is None,  "_cell_to_num(inf) = None")
check(P0A._cell_to_num(float('-inf')) is None, "_cell_to_num(-inf) = None")
check(P0A._cell_to_num(float('nan')) is None,  "_cell_to_num(nan) = None")
check(P0A._cell_to_num(2500) == 2500.0,        "_cell_to_num(2500) = 2500.0 (ค่าจริงคงอยู่)")
check(P0A._cell_to_num(1e30) == 1e30,          "_cell_to_num(1e30) = 1e30 (finite ผ่าน)")
check(P0A._cell_to_num('2,500') == 2500.0,     "_cell_to_num('2,500') = 2500.0 (comma คงอยู่)")

# ════════════════════════════════════════════════════════════════════════════
# Master-S1 — write_master_file ต้องไม่ทำลาย master จริง + backup ของผู้ใช้
# ════════════════════════════════════════════════════════════════════════════
print("\n[Master-S1] write_master_file: kill-before-atexit ต้องไม่ทำ master จริงหายถาวร")
REAL = {"บริษัท ก จำกัด": {"tax_id": "0000000000001"},
        "บริษัท ข จำกัด": {"tax_id": "0000000000002"}}


def _wj(p, obj):
    with open(p, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False)


def _rj(p):
    with open(p, "r", encoding="utf-8") as f:
        return json.load(f)


# (1) สถานการณ์ร้ายแรง: รอบก่อนถูก kill → live=stub, .user.bak=master จริง.
#     เรียก write_master_file ซ้ำ ต้อง "ไม่" ทับ .user.bak ด้วย stub.
mp = os.path.join(TMP, "master_companies.json")
bak = mp + ".user.bak"
_wj(bak, dict(REAL))                                   # backup = master จริง
_wj(mp, {**GS.MASTER, "_golden_stub": True})           # live = stub (ค้างจากรอบที่ถูก kill)
_quiet(lambda: GS.write_master_file(mp))
ok_bak = (os.path.exists(bak) and "บริษัท ก จำกัด" in _rj(bak)
          and not _rj(bak).get("_golden_stub"))
check(ok_bak, "live=stub + .user.bak=จริง → เรียกซ้ำไม่ทับ backup (master จริงรอด)")

# (2) [INF1 FIX แทน ADR-039 #2] live=master จริง → refresh .user.bak ให้สะท้อนของจริง "ก่อนรอบนี้"
#     เดิมพฤติกรรม "ไม่ทับ backup เก่า" ทำข้อมูลหาย (ดู case 2b). ใหม่: backup = live ปัจจุบัน (atomic).
_wj(mp, dict(REAL)); _wj(bak, {"บริษัท ค จำกัด": {"tax_id": "0000000000003"}})
_quiet(lambda: GS.write_master_file(mp))
check(os.path.exists(bak) and "บริษัท ก จำกัด" in _rj(bak)
      and "บริษัท ค จำกัด" not in _rj(bak),
      ".user.bak ถูก refresh จาก master จริงปัจจุบัน (ของเก่าค้างถูกแทนที่)")

# (2b) [INF1 ตัวจริง] kill → ผู้ใช้ใส่ master "ใหม่" → รัน golden → restore: master ใหม่ต้องรอด (ไม่ถูก backup เก่ากลืน)
mpd = os.path.join(TMP, "mloss.json"); bakd = mpd + ".user.bak"
_wj(mpd, {**GS.MASTER, "_golden_stub": True})          # live=stub (รอบก่อนถูก kill ก่อน atexit)
_wj(bakd, {"co_OLD": {"tax_id": "1"}})                 # .user.bak เก่าค้าง
_wj(mpd, {"co_NEW": {"tax_id": "2"}})                  # ผู้ใช้ใส่ master "ใหม่" (live=จริง)
_quiet(lambda: GS.write_master_file(mpd))              # รอบ golden ใหม่ (เขียน stub + refresh backup→co_NEW)
GS._restore_master_file(mpd, bakd)                     # จำลอง atexit
_final = _rj(mpd); _final.pop("_golden_stub", None)
check("co_NEW" in _final and "co_OLD" not in _final,
      "INF1: master ใหม่หลัง kill+edit ไม่ถูก backup เก่ากลืน (ไม่หายถาวร)")

# (3) วงจรปกติ (golden flow): real → write_master_file → restore คืน master จริงครบ
mp2 = os.path.join(TMP, "m2.json"); bak2 = mp2 + ".user.bak"
_wj(mp2, dict(REAL))
_quiet(lambda: GS.write_master_file(mp2))
check(GS._file_has_stub_marker(mp2), "ระหว่างรัน: ไฟล์เป็น golden stub (golden hash ไม่ขยับ)")
GS._restore_master_file(mp2, bak2)                     # จำลอง atexit
restored = _rj(mp2)
check("บริษัท ก จำกัด" in restored and "บริษัท ข จำกัด" in restored
      and not restored.get("_golden_stub"), "atexit restore คืน master จริงครบ 2 บริษัท")
check(not os.path.exists(bak2), "restore แล้ว .user.bak ถูก consume (atomic)")

# (4) atomic write: ได้ไฟล์ valid + ไม่เหลือ .tmp
mp3 = os.path.join(TMP, "m3.json")
GS._atomic_write_json(mp3, {"x": 1})
check(_rj(mp3) == {"x": 1} and not os.path.exists(mp3 + ".tmp"),
      "_atomic_write_json เขียน valid + ลบ .tmp")

# ════════════════════════════════════════════════════════════════════════════
# Master-M1/M2 — save_master atomic + .bak ห้ามหด
# ════════════════════════════════════════════════════════════════════════════
print("\n[Master-M1/M2] save_master: atomic write + .bak ไม่ถูก save ที่หดทับ")
_orig_cfg = mstr.CFG
try:
    smf = os.path.join(TMP, "save_master.json")
    mstr.CFG = {**dict(mstr.CFG), 'MASTER_FILE': smf}   # CFG จริงเป็น mappingproxy (อ่านอย่างเดียว) → patch ref ในโมดูล
    sbak = smf + ".bak"
    AB = {"บริษัท ก": {"tax_id": "1"}, "บริษัท ข": {"tax_id": "2"}}

    # M1: เขียน valid + ไม่เหลือ .tmp
    _quiet(lambda: mstr.save_master(AB))
    check(os.path.exists(smf) and _rj(smf) == AB and not os.path.exists(smf + ".tmp"),
          "save_master เขียน atomic (valid + ไม่เหลือ .tmp)")

    # M2: หดบริษัท (AB → A → {}) แล้ว .bak ต้องคงสภาพ "ครบสุด" (AB) ไว้กู้คืน
    _quiet(lambda: mstr.save_master({"บริษัท ก": {"tax_id": "1"}}))   # AB→A : .bak=AB
    _quiet(lambda: mstr.save_master({}))                              # A→{} : ต้องไม่ทับ .bak ด้วย A
    bak_after = _rj(sbak) if os.path.exists(sbak) else {}
    check("บริษัท ก" in bak_after and "บริษัท ข" in bak_after,
          f"save ที่หดไม่ทำ .bak สูญข้อมูล (.bak มี {len(bak_after)} บริษัท, ต้อง 2)")
finally:
    mstr.CFG = _orig_cfg

# ════════════════════════════════════════════════════════════════════════════
shutil.rmtree(TMP, ignore_errors=True)
print("\n" + "=" * 64)
print(f"BUGHUNT HARDENING: ผ่าน {PASS} / ล้มเหลว {len(FAIL)}")
print("=" * 64)
if FAIL:
    for f in FAIL:
        print(f"  • {f}")
    print("RESULT: ❌ พบ regression")
    sys.exit(1)
print("RESULT: ✅ การ์ด hardening ครบ ไม่มี regression")
sys.exit(0)
