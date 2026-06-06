# -*- coding: utf-8 -*-
# ruff: noqa: F401, F811  [F3] split-base re-export shim — re-export ขึ้น chain + local import ซ้ำได้
"""reporting_p2 — OBJ-MAINT layer 2 (extract คัดลอกเป๊ะ, byte-identical).
cascade toolkit จาก reporting_p1 (และชั้นล่างทั้งหมด)."""
from reporting_p1 import (   # [F3 de-star] explicit re-export shim (split-base chain; เดิม `import *`)
    ANALYTICS_CFG, APP_VERSION, Alignment, Border,
    COLORS, Counter, FIELD_CODES, Font,
    HTML, PYTHAINLP_AVAILABLE, PatternFill, REVIEW_CODES,
    RULES, SEVERITY_ORDER, Side, Workbook,
    _STY_BORDER, _STY_HEADER_ALIGN, _STY_HEADER_FILL, _STY_HEADER_FONT,
    _STY_SEV_FILLS, _STY_ZEBRA, _SYSTEM_ISSUES, _clean_company_label,
    _clean_period, _clean_sheet_allbills, _clean_sheet_duplicates, _clean_sheet_highrisk,
    _clean_sheet_items, _clean_sheet_summary, _get_p, _plotly_layout,
    _sty_cell_len, _sty_col_widths, _sty_find_cols, _sty_header,
    _sty_row, _sty_sev_row, _sty_zebra_row, _style_header,
    _style_one_sheet, _write_table, _xlsx_sheet_allbills, _xlsx_sheet_crossbill,
    _xlsx_sheet_dashboard, _xlsx_sheet_error_report, _xlsx_sheet_heatmap, _xlsx_sheet_highrisk,
    _xlsx_sheet_items, _xlsx_sheet_monthly, _xlsx_sheet_ranking, _xlsx_sheet_rules,
    _xlsx_sheet_summary, _xlsx_sheet_system_issues, bill_company_label, bottom_only,
    build_dashboard_figs, datetime, defaultdict, display,
    display_executive_dashboard, display_low_confidence_bills, export_excel, export_verification_to_excel,
    field_status, file_company_code, fill, generate_visual_dashboard,
    get_column_letter, go, issue_lane, kpi_cards,
    load_workbook, math, money_cards, np,
    os, pd, plt, predict_category,
    px, re, render_dashboard_html, risk_score_card,
    section_header, sort_bills_by_date, style_count, style_excel_report,
    style_severity, thin, traceback,
)  # noqa: F401  (re-export ขึ้น chain — หลายชื่อไม่ได้ใช้ภายในไฟล์นี้)
from config import CLEAN  # [F3] explicit — config ที่ reporting_p2 ใช้

def _clean_sheet_dashboard(wb, all_bills, summary, stats):
    """ชีต Dashboard (cards + financial band + charts). verbatim จาก _build_clean_report_impl."""
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    crit, err, warn, info, n_files, n_bills, n_clean, n_abn, n_items, n_issues, _ts, sum_sub, sum_vat, sum_tot, risk = (
        stats['crit'],
        stats['err'],
        stats['warn'],
        stats['info'],
        stats['n_files'],
        stats['n_bills'],
        stats['n_clean'],
        stats['n_abn'],
        stats['n_items'],
        stats['n_issues'],
        stats['_ts'],
        stats['sum_sub'],
        stats['sum_vat'],
        stats['sum_tot'],
        stats['risk'],
    )
    # ============ 1) DASHBOARD ============
    ws = wb.active
    ws.title = 'Dashboard'
    ws.sheet_view.showGridLines = False
    for col in range(1, 14):
        ws.column_dimensions[get_column_letter(col)].width = 12

    # แถบไตเติล
    ws.merge_cells('A1:M2')
    t = ws['A1']
    t.value = 'รายงานตรวจสอบบิล & เอกสารภาษี'
    t.font = Font(name=CLEAN['FONT'], bold=True, size=20, color='FFFFFF')
    t.alignment = Alignment(vertical='center', horizontal='left', indent=1)
    for col in range(1, 14):
        ws.cell(row=1, column=col).fill = fill(CLEAN['INK'])
        ws.cell(row=2, column=col).fill = fill(CLEAN['INK'])
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 14
    # เส้น accent ทองบางๆ (v6.4 แทนแดงหนา)
    ws.merge_cells('A3:M3')
    for col in range(1, 14):
        ws.cell(row=3, column=col).fill = fill(CLEAN['ACCENT'])
    ws.row_dimensions[3].height = 2
    ws['A4'] = f"สร้างเมื่อ {datetime.now().strftime('%d/%m/%Y %H:%M')}  •  {n_files} ไฟล์  •  {n_bills} บิล"
    ws['A4'].font = Font(name=CLEAN['FONT'], size=10, color=CLEAN['GRAY'])
    ws.merge_cells('A4:M4')

    # KPI cards
    cards = [
        ('บิลทั้งหมด', n_bills, CLEAN['INK']),
        ('บิลปกติ', n_clean, CLEAN['GREEN']),
        ('บิลผิดปกติ', n_abn, CLEAN['AMBER']),
        ('CRITICAL', crit, CLEAN['RED']),
        ('ERROR', err, CLEAN['AMBER']),
        ('RISK SCORE', risk, CLEAN['RED']),
    ]
    start = 6
    cw = 2  # card width in cols
    for idx, (label, val, color) in enumerate(cards):
        c0 = 1 + idx*cw
        c1 = c0 + cw - 1
        ws.merge_cells(start_row=start, start_column=c0, end_row=start, end_column=c1)
        ws.merge_cells(start_row=start+1, start_column=c0, end_row=start+1, end_column=c1)
        num = ws.cell(row=start, column=c0, value=val)
        num.font = Font(name=CLEAN['FONT'], bold=True, size=22, color=color)
        num.alignment = Alignment(vertical='center', horizontal='center')
        num.number_format = '#,##0'
        lab = ws.cell(row=start+1, column=c0, value=label)
        lab.font = Font(name=CLEAN['FONT'], size=10, color=CLEAN['GRAY'])
        lab.alignment = Alignment(vertical='top', horizontal='center')
        for rr in (start, start+1):
            for cc in range(c0, c1+1):
                ws.cell(row=rr, column=cc).fill = fill(CLEAN['CARD'])
    ws.row_dimensions[start].height = 34
    ws.row_dimensions[start+1].height = 18

    # การเงิน (row ถัดมา)
    fin = [('ยอดก่อน VAT', sum_sub), ('VAT รวม', sum_vat), ('ยอดสุทธิ', sum_tot)]
    fstart = start + 3
    fw = 4
    for idx, (label, val) in enumerate(fin):
        c0 = 1 + idx*fw; c1 = c0 + fw - 1
        ws.merge_cells(start_row=fstart, start_column=c0, end_row=fstart, end_column=c1)
        ws.merge_cells(start_row=fstart+1, start_column=c0, end_row=fstart+1, end_column=c1)
        num = ws.cell(row=fstart, column=c0, value=val)
        num.font = Font(name=CLEAN['FONT'], bold=True, size=16, color=CLEAN['INK'])
        num.alignment = Alignment(vertical='center', horizontal='center')
        num.number_format = '#,##0.00'
        lab = ws.cell(row=fstart+1, column=c0, value=label)
        lab.font = Font(name=CLEAN['FONT'], size=10, color=CLEAN['GRAY'])
        lab.alignment = Alignment(vertical='top', horizontal='center')
        for rr in (fstart, fstart+1):
            for cc in range(c0, c1+1):
                ws.cell(row=rr, column=cc).fill = fill(CLEAN['BAND'])
    ws.row_dimensions[fstart].height = 26

    # ---------- hidden sheet สำหรับ chart data ----------
    hd = wb.create_sheet('_chartdata')
    hd.sheet_state = 'hidden'
    # severity table
    hd['A1'] = 'ระดับ'; hd['B1'] = 'จำนวน'
    sev_data = [('Critical',crit),('Error',err),('Warning',warn),('Info',info)]
    for i,(k,v) in enumerate(sev_data, start=2):
        hd.cell(row=i, column=1, value=k); hd.cell(row=i, column=2, value=v)
    # top companies by issues
    comp_iss = sorted(
        [(s['key'], sum(len(b['issues']) for b in s['bills'])) for s in summary],
        key=lambda x: -x[1])
    comp_iss = [c for c in comp_iss if c[1] > 0][:10]
    hd['D1'] = 'บริษัท'; hd['E1'] = 'Issues'
    for i,(k,v) in enumerate(comp_iss, start=2):
        hd.cell(row=i, column=4, value=str(k)[:22]); hd.cell(row=i, column=5, value=v)
    # monthly net
    monthly = defaultdict(float)
    for b in all_bills:
        if b.get('iv_date'):
            monthly[_clean_period(b['iv_date'])] += (b['total'] or 0)
    mitems = sorted(monthly.items())
    hd['G1'] = 'งวด'; hd['H1'] = 'ยอดสุทธิ'
    for i,(k,v) in enumerate(mitems, start=2):
        hd.cell(row=i, column=7, value=k); hd.cell(row=i, column=8, value=v)

    # ---------- charts ----------
    chart_row = fstart + 3
    # Pie: severity — โชว์แค่ % บนชิ้น ใช้ legend บอกชื่อระดับ (กันป้ายรก)
    if (crit+err+warn+info) > 0:
        pie = PieChart(); pie.title = 'สัดส่วนความผิดปกติ'
        data = Reference(hd, min_col=2, min_row=1, max_row=1+len(sev_data))
        cats = Reference(hd, min_col=1, min_row=2, max_row=1+len(sev_data))
        pie.add_data(data, titles_from_data=True); pie.set_categories(cats)
        pie.height = 7; pie.width = 10
        dl = DataLabelList()
        dl.showPercent = True       # โชว์เฉพาะเปอร์เซ็นต์
        dl.showCatName = False      # ปิดชื่อหมวด (มีใน legend แล้ว)
        dl.showSerName = False      # ปิดชื่อ series ("จำนวน")
        dl.showVal = False          # ปิดตัวเลขดิบ
        dl.showLegendKey = False
        dl.showBubbleSize = False
        pie.dataLabels = dl
        ws.add_chart(pie, f'A{chart_row}')
    # Bar: top companies by issues
    if comp_iss:
        bar = BarChart(); bar.type='bar'; bar.title='Top บริษัทที่มีปัญหา'
        data = Reference(hd, min_col=5, min_row=1, max_row=1+len(comp_iss))
        cats = Reference(hd, min_col=4, min_row=2, max_row=1+len(comp_iss))
        bar.add_data(data, titles_from_data=True); bar.set_categories(cats)
        bar.height = 7; bar.width = 12; bar.legend = None
        try:
            from openpyxl.drawing.fill import PatternFillProperties, ColorChoice
            bar.series[0].graphicalProperties.solidFill = CLEAN['BLUE']
        except Exception: pass
        ws.add_chart(bar, f'G{chart_row}')
    # Bar: monthly net
    if len(mitems) >= 2:
        bar2 = BarChart(); bar2.type='col'; bar2.title='ยอดสุทธิรายงวด'
        data = Reference(hd, min_col=8, min_row=1, max_row=1+len(mitems))
        cats = Reference(hd, min_col=7, min_row=2, max_row=1+len(mitems))
        bar2.add_data(data, titles_from_data=True); bar2.set_categories(cats)
        bar2.height = 7; bar2.width = 22; bar2.legend = None
        try:
            bar2.series[0].graphicalProperties.solidFill = CLEAN['INK']
        except Exception: pass
        ws.add_chart(bar2, f'A{chart_row+15}')

def _clean_sheet_errors(wb, all_bills_s):
    """ชีต Error Report + ข้อควรตรวจสอบ + System Issues (แยกเลน FINDING/REVIEW; SystemIssues จาก _SYSTEM_ISSUES)."""
    # ============ 5) Error Report ============
    ws5 = wb.create_sheet('Error Report')
    ws5.sheet_view.showGridLines = False
    SEV_ORDER = {'CRITICAL':0,'ERROR':1,'WARNING':2,'INFO':3}
    e_rows = []
    for b in all_bills_s:
        for i in b['issues']:
            e_rows.append({'_lane':issue_lane(i),
                'Severity':i['severity'],'Code':i['code'],'หมวด':i['category'],
                'บริษัท':_clean_company_label(b),'เลขภาษี':(b.get('tax_id') or '-'),'งวดบัญชี':_clean_period(b['iv_date']),
                'วันที่':b['iv_date_str'],'IV':(b.get('iv_number_raw') or b.get('iv_number') or '-'),
                'ไฟล์':b['file'],'กฎ':i['name'],'รายละเอียด':i['detail']})
    e_rows.sort(key=lambda x: (SEV_ORDER.get(x['Severity'],9), x['Code']))
    # v8.5 [FIX-LANE]: แยก finding (ต้องแก้/ตรวจ) ออกจาก review (ข้อสังเกต) — ของทุกชิ้นยังอยู่ครบ
    find_rows = [{k:v for k,v in r.items() if k != '_lane'} for r in e_rows if r['_lane'] == 'FINDING']
    rev_rows  = [{k:v for k,v in r.items() if k != '_lane'} for r in e_rows if r['_lane'] == 'REVIEW']
    _f_crit = sum(1 for r in find_rows if r['Severity']=='CRITICAL')
    _f_err  = sum(1 for r in find_rows if r['Severity']=='ERROR')
    _f_warn = sum(1 for r in find_rows if r['Severity']=='WARNING')
    _write_table(ws5, find_rows,
                 band_title='รายงานข้อผิดพลาด — เฉพาะที่ต้องแก้/ตรวจ',
                 band_sub=f'{len(find_rows):,} ประเด็น · 🔴{_f_crit} 🟠{_f_err} 🟡{_f_warn}',
                 freeze_cols=2)
    
    # ============ 5b) ข้อควรตรวจสอบ (review/แนะนำ — ไม่ใช่ error) ============
    if rev_rows:
        ws5b = wb.create_sheet('ข้อควรตรวจสอบ')
        ws5b.sheet_view.showGridLines = False
        _write_table(ws5b, rev_rows,
                     band_title='ข้อควรตรวจสอบ — ข้อสังเกต/คำแนะนำ (ต้องคนยืนยัน)',
                     band_sub=f'{len(rev_rows):,} รายการ · ความเชื่อมั่นต่ำ/soft — ไม่นับเป็นข้อผิดพลาด')
    
    # ============ 5c) System Issues (กฎพัง/parse fail — ไม่ใช่ผลตรวจบิล) ============
    if _SYSTEM_ISSUES:
        ws5c = wb.create_sheet('System Issues')
        ws5c.sheet_view.showGridLines = False
        sys_rows = [{'Severity':s.get('severity','INFO'),'Code':s.get('code','SYS'),
                     'หมวด':s.get('category','SYSTEM'),'รายการ':s.get('name',''),
                     'ไฟล์':s.get('file',''),'ชีต':s.get('sheet',''),
                     'รายละเอียด':s.get('detail','')} for s in _SYSTEM_ISSUES]
        _write_table(ws5c, sys_rows,
                     band_title='System Issues — บันทึกข้อผิดพลาดของระบบ (ตามรอยได้)',
                     band_sub=f'{len(sys_rows):,} รายการ — กฎที่ทำงานไม่สำเร็จ/parse ไม่ผ่าน (ดู sidecar .jsonl)')

def _build_clean_report_impl(all_bills, summary, iv_issues, typos, filename_issues, path):
    """สร้างไฟล์ Excel รายงานสรุปแบบคลีน ผู้บริหารอ่านง่าย ชีตไม่เยอะ
    ชีต: Dashboard | Summary | ทุกบิล | รายการสินค้า | Error Report | High Risk | บิลซ้ำ
    """
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList

    # table/style helpers (fill/_style_header/_write_table) → module level (slice 3B)

    # ---------- เตรียมข้อมูล ----------
    all_bills_s = sorted(all_bills, key=_sk)
    crit = sum(1 for b in all_bills for i in b['issues'] if i['severity']=='CRITICAL')
    err  = sum(1 for b in all_bills for i in b['issues'] if i['severity']=='ERROR')
    warn = sum(1 for b in all_bills for i in b['issues'] if i['severity']=='WARNING')
    info = sum(1 for b in all_bills for i in b['issues'] if i['severity']=='INFO')
    n_files = len(set(b['file'] for b in all_bills))
    n_bills = len(all_bills)
    n_clean = sum(1 for b in all_bills if not b['issues'])
    n_abn   = n_bills - n_clean
    n_items = sum(len(b['items']) for b in all_bills)       # v6.3: ใช้ทำซับไตเติลแถบหัวเรื่อง
    n_issues = crit + err + warn + info
    _ts = datetime.now().strftime('%d/%m/%Y %H:%M')
    sum_sub = sum(b['subtotal'] or 0 for b in all_bills)
    sum_vat = sum(b['vat'] or 0 for b in all_bills)
    sum_tot = sum(b['total'] or 0 for b in all_bills)
    risk    = crit*10 + err*3 + warn

    wb = Workbook()

    # 1) Dashboard -> _clean_sheet_dashboard (slice 3E pass2)
    _clean_sheet_dashboard(wb, all_bills, summary, {
        'crit': crit,
        'err': err,
        'warn': warn,
        'info': info,
        'n_files': n_files,
        'n_bills': n_bills,
        'n_clean': n_clean,
        'n_abn': n_abn,
        'n_items': n_items,
        'n_issues': n_issues,
        '_ts': _ts,
        'sum_sub': sum_sub,
        'sum_vat': sum_vat,
        'sum_tot': sum_tot,
        'risk': risk,
    })

    # 2) Summary -> _clean_sheet_summary (slice 3E)
    _clean_sheet_summary(wb, all_bills, n_bills, n_files, _ts)

    # 3) ทุกบิล -> _clean_sheet_allbills (slice 3E)
    _clean_sheet_allbills(wb, all_bills_s, n_bills, n_files, _ts)

    # 4) รายการสินค้า -> _clean_sheet_items (slice 3E)
    _clean_sheet_items(wb, all_bills_s, n_items, n_files)

    # 5/5b/5c) Error Report + Review + System Issues -> _clean_sheet_errors (slice 3E)
    _clean_sheet_errors(wb, all_bills_s)

    # 6) High Risk -> _clean_sheet_highrisk (slice 3E)
    _clean_sheet_highrisk(wb, all_bills_s, _ts)

    # 7) บิลซ้ำ -> _clean_sheet_duplicates (slice 3E)
    _clean_sheet_duplicates(wb, all_bills)

    # v6.4: แท็บชีตสีตามประเภท (โทนหม่นสุภาพ) — ดูเป็นมืออาชีพ หาง่าย (ไม่กระทบข้อมูล)
    _tabcol = {'Dashboard':CLEAN['INK'], 'Summary':CLEAN['BLUE'], 'ทุกบิล':CLEAN['GRAY'],
               'รายการสินค้า':CLEAN['GREEN'], 'Error Report':CLEAN['RED'],
               'ข้อควรตรวจสอบ':'B7791F', 'System Issues':CLEAN['GRAY'],
               'High Risk':'8A1C13', 'บิลซ้ำ':CLEAN['AMBER']}
    for _nm, _col in _tabcol.items():
        if _nm in wb.sheetnames:
            try: wb[_nm].sheet_properties.tabColor = _col
            except Exception: pass

    wb.save(path)
    return True

def build_clean_report(all_bills, summary, iv_issues, typos, filename_issues, path):
    """v5.9 FIX-3: ตัวห่อกันพัง (safety wrapper)
    เดิม build_clean_report ไม่มี try/except ครอบ — ถ้า save พัง (ไฟล์ถูก lock/disk เต็ม/
    อักขระแปลก) โปรแกรมจะ crash หลัง parse เสร็จ = เสียงานทั้งหมด
    wrapper นี้จับ error แล้วคืน False อย่างสุภาพ (เหมือน export_excel) เพื่อให้ main()
    print แจ้งเตือนแทนการล้ม — logic ภายในไม่เปลี่ยนเลย

    [P2-NOTE SCALING] การเขียน Excel ใช้ pandas→openpyxl + per-cell styling เฉพาะแถวที่เน้น
      (efficient พอสำหรับหลักพัน-หมื่นบิล). ถ้าวันใดมีหลายหมื่น-แสนบิลแล้วช้า/แรมโต:
      ทางเลือกที่ปลอดภัย (เรียงตามความเสี่ยงน้อย→มาก):
        1) แบ่งรายงานเป็นหลายไฟล์ตามงวด/บริษัท (ไม่แตะ builder)
        2) เปลี่ยน engine เป็น xlsxwriter constant_memory mode (ต้องทดสอบ style ใหม่)
        3) batch styling: ใส่ style ทั้ง range แล้ว override เฉพาะ exception
      ⚠️ ทั้ง 2-3 ต้องผ่าน regression_oracle + ตรวจหน้าตารายงานด้วยตา (style ไม่ครอบใน oracle)
    """
    try:
        return _build_clean_report_impl(all_bills, summary, iv_issues, typos, filename_issues, path)
    except Exception as e:
        print(f'⚠️ สร้างรายงานคลีนล้มเหลว: {e}')
        traceback.print_exc()
        return False

def _sk(b):
    d = b.get('iv_date')
    if d is None: return (0, 0, 0)
    return (getattr(d,'year',0), getattr(d,'month',0), getattr(d,'day',1))


# OBJ-MAINT: auto-export ทุกชื่อ (รวม _ และ import) → from-import * cascade ครบ
__all__=[n for n in list(globals().keys()) if not n.startswith('__') and n!='annotations']
