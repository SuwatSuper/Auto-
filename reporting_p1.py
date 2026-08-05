# -*- coding: utf-8 -*-
# ruff: noqa: F401  [F3] split-base re-export shim — re-export ชื่อขึ้น chain (unused-internally โดยเจตนา)
"""reporting_p1 — OBJ-MAINT layer 1 (extract คัดลอกเป๊ะ, byte-identical).
cascade toolkit จาก reporting_p0 (และชั้นล่างทั้งหมด)."""
from reporting_p0 import (   # [F3 de-star] explicit re-export shim (split-base chain; เดิม `import *`)
    Alignment, Border, COLORS, Counter,
    FIELD_CODES, Font, HTML, PYTHAINLP_AVAILABLE,
    PatternFill, REVIEW_CODES, RULES, SEVERITY_ORDER,
    Side, Workbook, _STY_BORDER, _STY_HEADER_ALIGN,
    _STY_HEADER_FILL, _STY_HEADER_FONT, _STY_SEV_FILLS, _STY_ZEBRA,
    _SYSTEM_ISSUES, _get_p, _sty_cell_len, _sty_col_widths,
    _sty_find_cols, _sty_header, _sty_row, _sty_sev_row,
    _sty_zebra_row, _style_one_sheet, _xlsx_sheet_allbills, _xlsx_sheet_dashboard,
    _xlsx_sheet_error_report, _xlsx_sheet_highrisk, _xlsx_sheet_ranking, _xlsx_sheet_summary,
    _xlsx_sheet_system_issues, bill_company_label, datetime, defaultdict,
    display, display_executive_dashboard, display_low_confidence_bills, export_verification_to_excel,
    field_status, file_company_code, generate_visual_dashboard, get_column_letter,
    go, issue_lane, kpi_cards, load_workbook,
    math, money_cards, np, os,
    pd, plt, predict_category, px,
    re, risk_score_card, section_header, sort_bills_by_date,
    style_count, style_excel_report, style_severity, traceback,
)  # noqa: F401  (re-export ขึ้น chain — หลายชื่อไม่ได้ใช้ภายในไฟล์นี้)
from config import ANALYTICS_CFG, APP_VERSION, CLEAN, bottom_only, thin  # [F3] explicit — config ที่ reporting_p1 ใช้

try:
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE as _XL_ILLEGAL   # [H3] regex ของ openpyxl เอง
except Exception:                                                        # เผื่อ path เปลี่ยนข้ามเวอร์ชัน
    _XL_ILLEGAL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')

try:
    from openpyxl.cell.cell import KNOWN_TYPES as _XL_KNOWN_TYPES        # [ADR-168] ชนิดที่ openpyxl รับตรง ๆ
except Exception:                                                        # เผื่อ path เปลี่ยนข้ามเวอร์ชัน
    import datetime as _dt
    from decimal import Decimal as _Dec   # [ADR-170] fallback ต้องมี Decimal ให้ตรง openpyxl จริง
    _XL_KNOWN_TYPES = (int, float, str, bytes, bool, type(None), _Dec,
                       _dt.datetime, _dt.date, _dt.time, _dt.timedelta)


def _xl_safe(v):
    """[H3] ตัดอักขระควบคุมที่ openpyxl ปฏิเสธ (เช่น \\x07) ออกจากสตริงก่อนเขียนลงเซลล์.
    เดิม IllegalCharacterError จากเซลล์ที่ไม่ถูก sanitize (detail/name_raw/ชื่อไฟล์-ชีต) →
    build_clean_report คืน False → ผู้ใช้ "ไม่ได้รายงานเลย" ทั้งที่ตรวจเสร็จครบทั้งรอบ.
    [ADR-168/BUG-2 C4] ค่า non-scalar (list/dict/object หลุดมากับบิลเพี้ยน เช่น qty เป็น list) →
    openpyxl ValueError 'Cannot convert … to Excel' ล้มทั้ง workbook — แปลงเป็นสตริงแทน.
    ใช้ KNOWN_TYPES ของ openpyxl เองเป็นเกณฑ์ (รวม numpy scalar) → ค่าที่เคยเขียนได้ = เขียนเหมือนเดิมเป๊ะ
    (report-determinism ไม่ขยับ) ; เฉพาะค่าที่เดิม "ครัช" เท่านั้นที่กลายเป็นสตริง."""
    if isinstance(v, str):
        v = _XL_ILLEGAL.sub('', v)
        # [ADR-176 ← ปิด finding-2/ADR-175] formula injection: openpyxl ตีความสตริงขึ้นต้น '=' เป็น
        #   "สูตรจริง" (พิสูจน์: =HYPERLINK live). corpus จริงสแกน = 0 เซลล์ขึ้นต้น '=' → เกราะนี้
        #   zero-impact กับข้อมูลปัจจุบัน (REPORT_DET ไม่ขยับ) — คุ้มกันเฉพาะข้อมูลพิษอนาคต.
        return "'" + v if v.startswith('=') else v
    if isinstance(v, _XL_KNOWN_TYPES):   # ครอบ None ด้วย (type(None) อยู่ใน KNOWN_TYPES)
        return v
    return _XL_ILLEGAL.sub('', str(v))


def _fin(v):
    """[REP-C2 2026-06-20] ยอดเงินปลอดภัยสำหรับบวก/เขียนเซลล์ → คืน 0 เมื่อ None/NaN/±inf.
    เดิม report builder ใช้ `v or 0`: nan เป็น truthy → `nan or 0 == nan` ลามผ่าน sum() →
    openpyxl เขียน nan เป็น "เซลล์ว่าง" → Dashboard/Summary ยอดเงินหาย (เข้าใจผิดว่ายอด 0).
    M6 แก้ที่ analytics._num แล้ว แต่ report builder บวกยอดเอง → จุดนี้ยังโล่ง. ข้อมูลจริง
    subtotal/vat/total เป็น float จำกัด|None เสมอ → ค่าเท่าเดิม (report-determinism ไม่ขยับ)."""
    try:
        f = float(v)
    except (TypeError, ValueError):
        return 0
    return f if (f == f and f not in (float('inf'), float('-inf'))) else 0


def _xlsx_sheet_heatmap(writer, all_bills):
    """ชีต Error Heatmap (หมวด×severity)."""
    heatmap = defaultdict(lambda: defaultdict(int))
    for b in all_bills:
        for i in b['issues']: heatmap[i['category']][i['severity']] += 1
    if heatmap:
        rows = []
        for cat in heatmap:
            row = {'หมวด':cat}
            for sev in ['CRITICAL','ERROR','WARNING','INFO']: row[sev] = heatmap[cat][sev]
            row['รวม'] = sum(heatmap[cat].values())
            rows.append(row)
        pd.DataFrame(rows).sort_values('รวม', ascending=False).to_excel(writer, sheet_name='Error Heatmap', index=False)

def _xlsx_sheet_monthly(writer, all_bills):
    """ชีต Monthly Pattern (สรุปต่องวด)."""
    monthly = defaultdict(lambda: {'count':0,'subtotal':0,'vat':0,'total':0,'issues':0})
    for b in all_bills:
        if b['iv_date']:
            k = _get_p(b['iv_date'])
            monthly[k]['count'] += 1
            monthly[k]['subtotal'] += _fin(b['subtotal'])   # [REP-C2] NaN/inf-safe
            monthly[k]['vat'] += _fin(b['vat'])
            monthly[k]['total'] += _fin(b['total'])
            monthly[k]['issues'] += len(b['issues'])
    if monthly:
        pd.DataFrame([{'งวดบัญชี':k,**v} for k,v in sorted(monthly.items())]).to_excel(writer, sheet_name='Monthly Pattern', index=False)

def _xlsx_sheet_rules(writer):
    """ชีต Rules (รายการกฎทั้งหมด) — [A2] โชว์สถานะ 3 กลุ่ม + เหตุผล (กันโฆษณา 'active' หลอกตา)."""
    try:
        import code_registry as _reg
        _status, _reason = _reg.rule_status, _reg.rule_status_reason
    except Exception:                                  # advisory column — ขาด registry ไม่ทำรายงานล่ม
        _status = lambda c: ('active' if RULES[c].get('enabled', True) else 'disabled-by-design')
        _reason = lambda c: ''
    pd.DataFrame([{'Code':c,'Severity':r['severity'],'Category':r['category'],
        'Rule':r['name'],'Enabled':r['enabled'],
        'สถานะ':_status(c),'เหตุผล (ถ้าไม่ active)':_reason(c)} for c, r in RULES.items()
        ]).to_excel(writer, sheet_name='Rules', index=False)

def _xlsx_sheet_crossbill(writer, iv_issues, typos, filename_issues):
    """ชีต IV Cross-Bill / Product Typo / Filename (เขียนเมื่อมีข้อมูล)."""
    if iv_issues:
        pd.DataFrame(iv_issues).to_excel(writer, sheet_name='IV Cross-Bill', index=False)
    if typos:
        pd.DataFrame(typos).to_excel(writer, sheet_name='Product Typo', index=False)
    if filename_issues:
        pd.DataFrame(filename_issues).to_excel(writer, sheet_name='Filename', index=False)

def _xlsx_sheet_items(writer, all_bills_s):
    """ชีต รายการสินค้า (line items)."""
    items = []
    for b in all_bills_s:
        for it in b['items']:
            # [UPDATE] 4. ยัดคอลัมน์ 'งวดบัญชี' ลงชีต 'รายการสินค้า'
            items.append({'งวดบัญชี': _get_p(b['iv_date']), 'บริษัท':bill_company_label(b),
                'รหัสไฟล์':file_company_code(b['file']),
                'วันที่':b['iv_date_str'],'ไฟล์':b['file'],'ชีต':b['sheet'],
                'IV': (b.get('iv_number_raw') or b.get('iv_number') or '-'),
                'ลำดับ':it['seq'],'รายการ':it['name'],
                'จำนวน':it['qty'],'หน่วย':it['unit'],'ราคา':it['price'],'ยอด':it['amount']})
    if items:
        pd.DataFrame(items).to_excel(writer, sheet_name='รายการสินค้า', index=False)

def export_excel(all_bills, summary, iv_issues, typos, filename_issues, path):
    # decomposed (slice 3D): orchestrator คุม with/try; แต่ละชีตเป็น _xlsx_sheet_* — ลำดับ/เงื่อนไขเดิม
    try:
        # [A5-FIX] กัน issue ที่ขาดคีย์มาตรฐาน ทำทั้ง workbook ล่ม (เหมือน build_clean_report).
        #   setdefault = no-op กับ issue ปกติ → รายงาน/golden ไม่ขยับ.
        for _b in all_bills:
            # [ADR-173/BUGHUNT] twin ของ ADR-155/168 ฝั่ง full mode: issues=None ทำ len(b['issues'])
            #   ครัชในชีต ทุกบิล/dashboard → ทั้ง workbook หายเงียบ (คืน False).
            if not isinstance(_b.get('issues'), list): _b['issues'] = []
            for _i in _b['issues']:
                # [M4] coerce None เช่นกัน (detail=None → None[:80] ครัช ทั้งที่ setdefault ผ่าน)
                _i['code'] = _i.get('code') or 'UNKNOWN'; _i['severity'] = _i.get('severity') or 'INFO'
                _i['category'] = _i.get('category') or '-'; _i['detail'] = _i.get('detail') or ''
                _i['name'] = _i.get('name') or _i.get('code') or ''
                # [ADR-173] twin ของ ADR-168 C1: ค่าชนิดผิดแต่ truthy (int code) → str() (no-op ปกติ)
                for _k in ('code', 'severity', 'category', 'detail', 'name'):
                    if not isinstance(_i[_k], str): _i[_k] = str(_i[_k])
        all_bills_s = sort_bills_by_date(all_bills)
        with pd.ExcelWriter(path, engine='openpyxl') as writer:
            _xlsx_sheet_dashboard(writer, all_bills)
            _xlsx_sheet_summary(writer, summary)
            _xlsx_sheet_allbills(writer, all_bills_s)
            _xlsx_sheet_highrisk(writer, all_bills_s)
            _xlsx_sheet_ranking(writer, summary)
            _xlsx_sheet_error_report(writer, all_bills_s)
            _xlsx_sheet_system_issues(writer)
            _xlsx_sheet_heatmap(writer, all_bills)
            _xlsx_sheet_monthly(writer, all_bills)
            _xlsx_sheet_rules(writer)
            _xlsx_sheet_crossbill(writer, iv_issues, typos, filename_issues)
            _xlsx_sheet_items(writer, all_bills_s)
        style_excel_report(path)
        return True
    except Exception as e:
        print(f'⚠️ Export ล้มเหลว: {e}')
        traceback.print_exc()
        return False

def _plotly_layout(title):
    return dict(
        title=dict(text=title, font=dict(size=18, color='#1E40AF')),
        font=dict(family='sans-serif', size=12),
        plot_bgcolor='white', paper_bgcolor='white',
        margin=dict(t=60, b=40, l=60, r=40),
    )

def build_dashboard_figs(sheets, df_clustered):
    try:
        import plotly.express as px
        import plotly.graph_objects as go
    except ImportError:
        print('⚠️ Plotly ยังไม่ติดตั้ง')
        return []
    figs = []
    if 'ยอด' in df_clustered.columns and '_category' in df_clustered.columns:
        d = df_clustered.copy()
        d['_amount'] = pd.to_numeric(d['ยอด'], errors='coerce').fillna(0)
        d = d[d['_amount'] > 0]
        if len(d) > 0:
            try:
                fig = px.treemap(d, path=['_category','_cluster_label'], values='_amount',
                                title='🗺️ Treemap: หมวด → กลุ่มสินค้า → ยอดรวม',
                                color='_amount', color_continuous_scale='Blues')
                fig.update_layout(**_plotly_layout(''), height=550)
                figs.append(('treemap_category', fig))
            except Exception as e: print(f'   Treemap fail: {e}')
    if 'ราคา' in df_clustered.columns and 'จำนวน' in df_clustered.columns:
        d = df_clustered.copy()
        d['_price'] = pd.to_numeric(d['ราคา'], errors='coerce')
        d['_qty'] = pd.to_numeric(d['จำนวน'], errors='coerce')
        d = d[(d['_price'] > 0) & (d['_qty'] > 0)]
        if len(d) > 0:
            try:
                hover_cols = [c for c in ['รายการ','หน่วย','IV','_cluster_label'] if c in d.columns]
                fig = px.scatter(d, x='_qty', y='_price', color='_category',
                                hover_data=hover_cols, log_x=True, log_y=True,
                                title='💰 Scatter: ราคา × จำนวน (log scale, color=หมวด)',
                                labels={'_qty':'จำนวน','_price':'ราคา/หน่วย'})
                if '_price_outlier' in d.columns:
                    out = d[d['_price_outlier']]
                    if len(out) > 0:
                        fig.add_trace(go.Scatter(x=out['_qty'], y=out['_price'],
                                                mode='markers', name='Price Outlier',
                                                marker=dict(size=14, symbol='x', color='red',
                                                          line=dict(width=2))))
                fig.update_layout(**_plotly_layout(''), height=500)
                figs.append(('scatter_price_qty', fig))
            except Exception as e: print(f'   Scatter fail: {e}')
    if 'ราคา' in df_clustered.columns and '_category' in df_clustered.columns:
        d = df_clustered.copy()
        d['_price'] = pd.to_numeric(d['ราคา'], errors='coerce')
        d = d[d['_price'] > 0]
        if len(d) > 0:
            try:
                fig = px.box(d, x='_category', y='_price', color='_category',
                            title='📦 Box Plot: การกระจายตัวของราคาในแต่ละหมวด',
                            log_y=True)
                fig.update_layout(**_plotly_layout(''), showlegend=False, height=500)
                figs.append(('box_price_category', fig))
            except Exception as e: print(f'   Box fail: {e}')
    bills = sheets.get('ทุกบิล')
    items = sheets.get('รายการสินค้า')
    if items is not None and 'รายการ' in items.columns and 'ไฟล์' in items.columns:
        d = items.copy()
        d['_cat'] = d['รายการ'].fillna('').apply(predict_category)
        # [M1] กันคอลัมน์ 'หน่วย' หาย: d.get('หน่วย','') คืน scalar '' → ''.fillna() = AttributeError (crash ทั้ง dashboard)
        if 'หน่วย' in d.columns:
            d['_unit'] = d['หน่วย'].fillna('-')
        else:
            d['_unit'] = '-'
        d['_unit'] = d['_unit'].replace('', '-')
        d['_count'] = 1
        agg = d.groupby(['ไฟล์','_cat','_unit'])['_count'].sum().reset_index()
        agg = agg[agg['_count'] > 0]
        if len(agg) > 0:
            try:
                fig = px.sunburst(agg, path=['ไฟล์','_cat','_unit'], values='_count',
                                 title='☀️ Sunburst: ไฟล์ → หมวด → หน่วย')
                fig.update_layout(**_plotly_layout(''), height=600)
                figs.append(('sunburst_file_cat_unit', fig))
            except Exception as e: print(f'   Sunburst fail: {e}')
    verif = sheets.get('Product Verification')
    if verif is not None and 'Final State' in verif.columns and 'หมวดที่คาด' in verif.columns:
        try:
            pivot = verif.pivot_table(index='หมวดที่คาด', columns='Final State',
                                      values='ชื่อสินค้าเดิม', aggfunc='count', fill_value=0)
            fig = px.imshow(pivot.values, x=pivot.columns, y=pivot.index,
                           color_continuous_scale='RdYlGn_r', text_auto=True,
                           title='🔥 Heatmap: หมวด × Verification State',
                           aspect='auto')
            fig.update_layout(**_plotly_layout(''), height=400)
            figs.append(('heatmap_verif_category', fig))
        except Exception as e: print(f'   Heatmap fail: {e}')
    if '_cluster_id' in df_clustered.columns:
        # [M2] กัน KeyError: False — ถ้าทั้งสองคอลัมน์หาย df.get(...) คืน scalar False → False|False=False → df[False] crash
        _po = df_clustered['_price_outlier'] if '_price_outlier' in df_clustered.columns else False
        _ur = df_clustered['_unit_rare'] if '_unit_rare' in df_clustered.columns else False
        _mask = _po | _ur
        d = df_clustered[_mask] if hasattr(_mask, '__len__') else df_clustered.iloc[0:0]
        if len(d) > 0:
            try:
                top = d.groupby('_cluster_label').size().reset_index(name='count')
                top = top.sort_values('count', ascending=False).head(ANALYTICS_CFG['TOP_N_DISPLAY'])
                fig = px.bar(top, x='count', y='_cluster_label', orientation='h',
                            title=f'⚠️ Top {len(top)} กลุ่มสินค้าที่มี Anomaly',
                            color='count', color_continuous_scale='Reds')
                fig.update_layout(**_plotly_layout(''), height=max(400, len(top)*20))
                fig.update_yaxes(autorange='reversed')
                figs.append(('bar_top_anomalies', fig))
            except Exception as e: print(f'   Bar fail: {e}')
    errors = sheets.get('Error Report')
    if errors is not None and 'Severity' in errors.columns:
        try:
            sev = errors['Severity'].value_counts().reset_index()
            sev.columns = ['Severity','Count']
            color_map = {'CRITICAL':'#DC2626','ERROR':'#EA580C',
                        'WARNING':'#EAB308','INFO':'#0EA5E9'}
            fig = px.pie(sev, names='Severity', values='Count',
                        title='🚨 สัดส่วน Severity ของ Issues ทั้งหมด',
                        color='Severity', color_discrete_map=color_map, hole=0.4)
            fig.update_layout(**_plotly_layout(''), height=400)
            figs.append(('pie_severity', fig))
        except Exception as e: print(f'   Pie fail: {e}')
    if bills is not None and 'วันที่' in bills.columns:
        try:
            d = bills.copy()
            d['_date'] = pd.to_datetime(d['วันที่'], errors='coerce', dayfirst=True)
            d = d.dropna(subset=['_date'])
            if len(d) > 0:
                d['_issues'] = pd.to_numeric(d.get('Issues', 0), errors='coerce').fillna(0)
                d['_amount'] = pd.to_numeric(d.get('ยอดก่อน VAT', 0), errors='coerce').fillna(0)
                hover_cols = [c for c in ['ไฟล์','IV','บริษัท'] if c in d.columns]
                fig = px.scatter(d, x='_date', y='_amount', size='_amount',
                                color='_issues', hover_data=hover_cols,
                                color_continuous_scale='Reds',
                                title='📅 Timeline: บิลตามเวลา (size=ยอด, color=Issues)')
                fig.update_layout(**_plotly_layout(''), height=450)
                figs.append(('timeline_bills', fig))
        except Exception as e: print(f'   Timeline fail: {e}')
    return figs

def render_dashboard_html(figs, output_path, stats=None):
    try:
        import plotly.io as pio
    except ImportError:
        print('⚠️ Plotly ไม่พร้อม')
        return False
    if not figs:
        print('⚠️ ไม่มี chart ให้ render')
        return False
    include_js = ANALYTICS_CFG['INCLUDE_PLOTLY_JS']
    stats = stats or {}
    parts = ['''<!DOCTYPE html>
<html lang="th"><head><meta charset="utf-8">
<title>Anomaly Dashboard — Audit v''' + APP_VERSION + '''</title>
<style>
body{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',sans-serif;
     margin:0;padding:24px;background:#F3F4F6;color:#111827;}
.header{background:linear-gradient(135deg,#0F172A,#1E40AF);color:white;
        padding:30px;border-radius:12px;margin-bottom:20px;
        box-shadow:0 10px 25px rgba(0,0,0,0.2);}
.header h1{margin:0;font-size:28px;}
.header p{margin:8px 0 0 0;opacity:0.9;font-size:14px;}
.kpi-grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(200px,1fr));
          gap:16px;margin-bottom:24px;}
.kpi{background:white;padding:18px;border-radius:10px;
     box-shadow:0 2px 6px rgba(0,0,0,0.06);border-left:5px solid #3B82F6;}
.kpi-label{color:#6B7280;font-size:12px;font-weight:500;}
.kpi-value{color:#1E40AF;font-size:28px;font-weight:700;line-height:1.1;}
.chart{background:white;padding:16px;border-radius:10px;margin-bottom:20px;
       box-shadow:0 2px 6px rgba(0,0,0,0.06);}
.footer{text-align:center;color:#6B7280;font-size:12px;margin-top:30px;}
</style></head><body>
<div class="header">
  <h1>📊 Anomaly Dashboard</h1>
  <p>Pattern Clustering + Plotly Interactive Charts • Audit System v''' + APP_VERSION + '''</p>
  <p>Generated: ''' + datetime.now().strftime('%Y-%m-%d %H:%M') + '''</p>
</div>''']
    if stats:
        parts.append('<div class="kpi-grid">')
        for label, value in stats.items():
            parts.append(f'<div class="kpi"><div class="kpi-label">{label}</div>'
                         f'<div class="kpi-value">{value}</div></div>')
        parts.append('</div>')
    for i, (name, fig) in enumerate(figs):
        html_fig = pio.to_html(fig, include_plotlyjs=(include_js if i == 0 else False),
                              full_html=False, div_id=f'chart_{name}')
        parts.append(f'<div class="chart">{html_fig}</div>')
    parts.append(f'<div class="footer">{len(figs)} charts • Method: {ANALYTICS_CFG["CLUSTERING_METHOD"]}'
                 f' • Threshold: {ANALYTICS_CFG["FUZZY_CLUSTER_THRESHOLD"]}%</div>')
    parts.append('</body></html>')
    try:
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(parts))
        return True
    except Exception as e:
        print(f'⚠️ บันทึก HTML ไม่ได้: {e}')
        return False

def _clean_period(d):
    if not d: return '-'
    # [ADR-168/BUG-2 C2] iv_date ชนิดผิดแต่ truthy (str/float serial จากบิลเพี้ยน) → เดิม d.year ครัช
    #   AttributeError ล้มทั้ง workbook (เส้นคลีนเรียก 6 จุด) — การ์ดแบบเดียวกับ _sk (getattr).
    #   [ADR-170] int() ครอบด้วย: pd.NaT truthy + .year=nan ทะลุเช็ค None แล้วครัชที่ f-string.
    #   วันที่ปกติ (datetime.date) ได้ผลเท่าเดิมเป๊ะ.
    try: yr = int(getattr(d, 'year', None)); mo = int(getattr(d, 'month', None))
    except (TypeError, ValueError): return '-'
    be = yr + 543 if yr < 2500 else yr
    return f"{(be % 100):02d}.{mo:02d}"

def _clean_company_label(b):
    mk = b.get('master_key')
    # [ADR-168/BUG-2 C3] str() กัน master_key/company ชนิดผิด (float/int) → เดิมหลุดไป sorted() ใน
    #   _clean_sheet_summary เทียบ str กับ float → TypeError ล้มทั้ง workbook. ค่าปกติ (str) = no-op.
    if mk and mk != '(ไม่พบใน master)': return str(mk)
    comp = str(b.get('company') or '').strip()
    if comp: return comp
    base = os.path.basename(str(b.get('file','')))
    m = re.match(r'\s*([A-Za-zก-๙]+)', base)
    return m.group(1) if m else '[ตรวจสอบเอง]'

def fill(hexc): return PatternFill('solid', fgColor=hexc)

def _style_header(ws, row, ncols, ink=None):
    ink = ink or CLEAN['INK']
    accent = Side(style='thin', color=CLEAN['ACCENT'])   # v6.4: เส้นทองบางๆ ใต้หัวตาราง (แทนแดงหนา)
    for c in range(1, ncols+1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill(ink)
        cell.font = Font(name=CLEAN['FONT'], bold=True, color='FFFFFF', size=10)
        cell.alignment = Alignment(vertical='center', horizontal='left', wrap_text=False)
        cell.border = Border(bottom=accent)
    ws.row_dimensions[row].height = 26

def _write_table(ws, rows, start_row=1, money_cols=(), int_cols=(), title=None,
                 total_cols=(), band_title=None, band_sub=None, freeze_cols=0):
    """เขียนตารางจาก list-of-dict แบบคลีน คืนเลขแถวถัดไป
    v6.2: total_cols=(...) → เพิ่มแถว 'รวมทั้งหมด' (ผลรวม) ใต้ตาราง แบบเรียบหรู
          (วางนอกช่วง auto_filter เพื่อไม่ให้ถูกกรอง/ซ่อน)
    v6.3: band_title/band_sub → แถบหัวเรื่องสไตล์ผู้บริหาร (ดำ+เส้นแดง) เหนือตาราง
          freeze_cols=N → ตรึง N คอลัมน์แรก (เลื่อนขวาแล้วยังเห็น บริษัท/ไฟล์)"""
    r = start_row
    if title:
        ws.cell(row=r, column=1, value=_xl_safe(title)).font = Font(name=CLEAN['FONT'], bold=True, size=13, color=CLEAN['INK'])
        ws.row_dimensions[r].height = 24
        r += 2
    if not rows:
        ws.cell(row=r, column=1, value='— ไม่มีข้อมูล —').font = Font(name=CLEAN['FONT'], italic=True, color=CLEAN['GRAY'])
        return r + 1
    headers = list(rows[0].keys())
    ncols = len(headers)
    # v6.4: แถบหัวเรื่องเหนือตาราง (executive band) — ถ่านสุภาพ + ซับไตเติลเทาคูล + เส้นทองบางๆ
    if band_title:
        last = get_column_letter(ncols)
        ws.merge_cells(f"A{r}:{last}{r}")
        ws.merge_cells(f"A{r+1}:{last}{r+1}")
        tc = ws.cell(row=r, column=1, value=_xl_safe(band_title))
        tc.font = Font(name=CLEAN['FONT'], bold=True, size=15, color='FFFFFF')
        tc.alignment = Alignment(vertical='center', horizontal='left', indent=1)
        sc = ws.cell(row=r+1, column=1, value=_xl_safe(band_sub or ''))
        sc.font = Font(name=CLEAN['FONT'], size=9, color='AEB6C2')
        sc.alignment = Alignment(vertical='center', horizontal='left', indent=1)
        for col in range(1, ncols+1):
            ws.cell(row=r, column=col).fill = fill(CLEAN['INK'])
            ws.cell(row=r+1, column=col).fill = fill(CLEAN['INK'])
        ws.row_dimensions[r].height = 26
        ws.row_dimensions[r+1].height = 15
        ws.merge_cells(f"A{r+2}:{last}{r+2}")
        for col in range(1, ncols+1):
            ws.cell(row=r+2, column=col).fill = fill(CLEAN['ACCENT'])   # v6.4: ทองบางๆ แทนแดงหนา
        ws.row_dimensions[r+2].height = 2
        ws.row_dimensions[r+3].height = 6   # spacer
        r += 4
    for j, h in enumerate(headers, 1):
        ws.cell(row=r, column=j, value=h)
    _style_header(ws, r, len(headers))
    _fz_col = (int(freeze_cols) + 1) if freeze_cols else 1
    ws.freeze_panes = ws.cell(row=r+1, column=_fz_col)
    hdr_row = r
    r += 1
    for i, row in enumerate(rows):
        band = (i % 2 == 1)
        for j, h in enumerate(headers, 1):
            val = row.get(h)
            cell = ws.cell(row=r, column=j, value=_xl_safe(val))   # [H3] กัน IllegalCharacterError
            cell.font = Font(name=CLEAN['FONT'], size=10, color=CLEAN['INK2'])
            cell.alignment = Alignment(vertical='center', horizontal='left')
            cell.border = bottom_only
            if band: cell.fill = fill(CLEAN['BAND'])
            if h in money_cols:
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(vertical='center', horizontal='right')
            elif h in int_cols:
                cell.number_format = '#,##0'
                cell.alignment = Alignment(vertical='center', horizontal='right')
            elif h == 'จำนวน':
                # v6.2 polish: จำนวนชิดขวา + คั่นหลักพัน (โชว์ทศนิยมเฉพาะเมื่อมี)
                cell.number_format = '#,##0.###'
                cell.alignment = Alignment(vertical='center', horizontal='right')
            elif h in ('งวดบัญชี', 'งวด', 'IV', 'เลขภาษี'):
                cell.number_format = '@'   # เก็บเป็นข้อความ (กัน 69.10 → 69.1 / กันเลขภาษี 13 หลักหาย 0 นำหน้า)
            # ระบายสี severity
            if h == 'Severity':
                cmap = {'CRITICAL':CLEAN['RED'],'ERROR':CLEAN['AMBER'],
                        'WARNING':'B7791F','INFO':CLEAN['BLUE']}
                if val in cmap:
                    cell.font = Font(name=CLEAN['FONT'], size=10, bold=True, color=cmap[val])
        r += 1
    # auto width
    for j, h in enumerate(headers, 1):
        maxlen = len(str(h))
        for row in rows[:200]:
            v = row.get(h)
            if v is not None:
                maxlen = max(maxlen, len(str(v)))
        ws.column_dimensions[get_column_letter(j)].width = min(max(maxlen + 3, 10), 55)
    # ✅ เพิ่มฟิลเตอร์ (ปุ่มลูกศรกรอง/เรียงข้อมูล) บนหัวตาราง
    last_col = get_column_letter(len(headers))
    last_row = r - 1
    ws.auto_filter.ref = f"A{hdr_row}:{last_col}{last_row}"
    # v6.2 polish: แถวสรุป 'รวมทั้งหมด' ใต้ตาราง (นอกช่วง filter) — ดูเรียบหรู มืออาชีพ
    if total_cols and rows:
        top = Side(style='thin', color=CLEAN['INK'])
        for j, h in enumerate(headers, 1):
            cell = ws.cell(row=r, column=j)
            cell.border = Border(top=top)
            cell.fill = fill(CLEAN['CARD'])
            if j == 1:
                cell.value = f'รวมทั้งหมด ({len(rows):,} รายการ)'
                cell.font = Font(name=CLEAN['FONT'], bold=True, size=10, color=CLEAN['INK'])
                cell.alignment = Alignment(vertical='center', horizontal='left')
            elif h in total_cols:
                s = 0.0
                for row in rows:
                    v = row.get(h)
                    if isinstance(v, (int, float)): s += float(v)
                cell.value = s
                cell.number_format = '#,##0.00'
                cell.font = Font(name=CLEAN['FONT'], bold=True, size=10, color=CLEAN['INK'])
                cell.alignment = Alignment(vertical='center', horizontal='right')
            else:
                cell.font = Font(name=CLEAN['FONT'], bold=True, size=10, color=CLEAN['INK'])
        ws.row_dimensions[r].height = 22
        r += 1
    return r + 1  # v5.9 FIX-3: คืนค่าเดียว (เดิมคืน tuple ซึ่งไม่มีที่ไหนใช้ hdr_row)

def _clean_sheet_allbills(wb, all_bills_s, n_bills, n_files, _ts):
    """ชีต ทุกบิล."""
    # ============ 3) ทุกบิล ============
    ws3 = wb.create_sheet('ทุกบิล')
    ws3.sheet_view.showGridLines = False
    b_rows = [{'งวดบัญชี':_clean_period(b['iv_date']),'วันที่':b['iv_date_str'],'บริษัท':_clean_company_label(b),'เลขภาษี':(b.get('tax_id') or '-'),
               'IV':(b.get('iv_number_raw') or b.get('iv_number') or '-'),'ไฟล์':b['file'],
               'ยอดก่อน VAT':b['subtotal'],'VAT':b['vat'],'ยอดสุทธิ':b['total'],
               'รายการ':len(b['items']),'Issues':len(b['issues']),
               'Critical':sum(1 for i in b['issues'] if i['severity']=='CRITICAL'),
               'Error':sum(1 for i in b['issues'] if i['severity']=='ERROR')}
              for b in all_bills_s]
    _write_table(ws3, b_rows, money_cols=('ยอดก่อน VAT','VAT','ยอดสุทธิ'),
                 int_cols=('รายการ','Issues','Critical','Error'),
                 band_title='ทุกบิล — รายการเอกสารทั้งหมด',
                 band_sub=f'{n_bills:,} บิล · {n_files:,} ไฟล์ · ปรับปรุง {_ts}',
                 freeze_cols=3)

def _clean_sheet_highrisk(wb, all_bills_s, _ts):
    """ชีต High Risk (บิลที่มี CRITICAL)."""
    # ============ 6) High Risk (เฉพาะ critical) ============
    hr = [b for b in all_bills_s if any(i['severity']=='CRITICAL' for i in b['issues'])]
    if hr:
        ws6 = wb.create_sheet('High Risk')
        ws6.sheet_view.showGridLines = False
        h_rows = [{'งวดบัญชี':_clean_period(b['iv_date']),'วันที่':b['iv_date_str'],'บริษัท':_clean_company_label(b),
                   'IV':(b.get('iv_number_raw') or b.get('iv_number') or '-'),'ไฟล์':b['file'],'ยอด':_fin(b['subtotal']),
                   'Critical':sum(1 for i in b['issues'] if i['severity']=='CRITICAL'),
                   'หมายเหตุ':' | '.join(f"[{i['code']}] {(i.get('detail') or '')[:80]}" for i in b['issues'] if i['severity']=='CRITICAL')}
                  for b in hr]
        _write_table(ws6, h_rows, money_cols=('ยอด',), int_cols=('Critical',),
                     band_title='ความเสี่ยงสูง — บิลที่มี CRITICAL',
                     band_sub=f'{len(hr):,} บิลต้องตรวจสอบเร่งด่วน · ปรับปรุง {_ts}')

def _clean_sheet_items(wb, all_bills_s, n_items, n_files):
    """ชีต รายการสินค้า (line items)."""
    # ============ 4) รายการสินค้า ============
    ws4 = wb.create_sheet('รายการสินค้า')
    ws4.sheet_view.showGridLines = False
    i_rows = []
    for b in all_bills_s:
        for it in b['items']:
            i_rows.append({'งวดบัญชี':_clean_period(b['iv_date']),'บริษัท':_clean_company_label(b),
                'ไฟล์':b.get('file',''),'ชีต':str(b.get('sheet','')),
                'วันที่':b['iv_date_str'],'IV':(b.get('iv_number_raw') or b.get('iv_number') or '-'),
                'ลำดับ':it['seq'],'รายการ':it['name'],'จำนวน':it['qty'],
                'หน่วย':it['unit'],'ราคา':it['price'],'ยอด':it['amount']})
    _write_table(ws4, i_rows, money_cols=('ราคา','ยอด'), int_cols=('ลำดับ',),
                 total_cols=('ยอด',),
                 band_title='รายการสินค้า — รายละเอียดทุกบรรทัด',
                 band_sub=f'{n_items:,} รายการ · {n_files:,} ไฟล์ · ตรึงคอลัมน์ระบุไฟล์ไว้ซ้ายมือ',
                 freeze_cols=6)

def _clean_sheet_summary(wb, all_bills, n_bills, n_files, _ts):
    """ชีต Summary (บริษัท×งวดบัญชี)."""
    # ============ 2) SUMMARY — แยกตาม บริษัท × งวดบัญชี (เห็นครบทุกเดือน) ============
    ws2 = wb.create_sheet('Summary')
    ws2.sheet_view.showGridLines = False
    agg = defaultdict(lambda: {'bills':0,'sub':0.0,'vat':0.0,'tot':0.0,'iss':0,'crit':0,'tax':''})
    for b in all_bills:
        comp = _clean_company_label(b)
        per  = _clean_period(b['iv_date'])
        a = agg[(comp, per)]
        if not a['tax'] and (b.get('tax_id') or '').strip():   # v9.1: เก็บเลขภาษี (ดึงจากบิลจริง) ตัวแรกที่เจอของกลุ่ม
            a['tax'] = b['tax_id']
        a['bills'] += 1
        a['sub']   += _fin(b['subtotal'])   # [REP-C2] NaN/inf-safe (เดิม `or 0` ปล่อย nan ลามทำยอดว่าง)
        a['vat']   += _fin(b['vat'])
        a['tot']   += _fin(b['total'])
        a['iss']   += len(b['issues'])
        a['crit']  += sum(1 for i in b['issues'] if i['severity']=='CRITICAL')
    s_rows = [{'บริษัท':comp, 'เลขภาษี':(a['tax'] or '-'), 'งวดบัญชี':per, 'จำนวนบิล':a['bills'],
               'ยอดก่อน VAT':round(a['sub'],2), 'VAT':round(a['vat'],2),
               'ยอดสุทธิ':round(a['tot'],2), 'Issues':a['iss'], 'Critical':a['crit']}
              for (comp, per), a in sorted(agg.items(), key=lambda kv: (kv[0][0], kv[0][1]))]
    _write_table(ws2, s_rows, money_cols=('ยอดก่อน VAT','VAT','ยอดสุทธิ'),
                 int_cols=('จำนวนบิล','Issues','Critical'),
                 band_title='สรุปภาพรวม — ตามบริษัท × งวดบัญชี',
                 band_sub=f'{n_bills:,} บิล · {n_files:,} ไฟล์ · ปรับปรุง {_ts}')

def _clean_sheet_duplicates(wb, all_bills):
    """ชีต บิลซ้ำ (IV+ยอด+บริษัท ตรงกัน)."""
    # ============ 7) บิลซ้ำ ============
    seen = defaultdict(list)
    for b in all_bills:
        key = (b.get('iv_number') or '-', round(_fin(b['total']),2), _clean_company_label(b))
        seen[key].append(b)
    dups = [(k,v) for k,v in seen.items() if len(v) > 1 and k[0] != '-']
    if dups:
        ws7 = wb.create_sheet('บิลซ้ำ')
        ws7.sheet_view.showGridLines = False
        d_rows = []
        for (iv, amt, comp), bs in dups:
            iv_disp = bs[0].get('iv_number_raw') or iv  # [DISPLAY-FAITHFUL] เลขดิบตามใบ (key ใช้ norm สำหรับ dedup แต่ช่องโชว์ต้องตรงใบ)
            d_rows.append({'IV':iv_disp,'บริษัท':comp,'ยอดสุทธิ':amt,'พบซ้ำ':len(bs),
                           'ไฟล์':' | '.join(sorted(set(b['file'] for b in bs)))})
        _write_table(ws7, d_rows, money_cols=('ยอดสุทธิ',), int_cols=('พบซ้ำ',),
                     band_title='บิลซ้ำ — ตรวจพบความเป็นไปได้ที่ซ้ำกัน',
                     band_sub=f'{len(dups):,} กลุ่ม (IV + ยอด + บริษัท ตรงกัน)')


# OBJ-MAINT: auto-export ทุกชื่อ (รวม _ และ import) → from-import * cascade ครบ
__all__=[n for n in list(globals().keys()) if not n.startswith('__') and n!='annotations']
